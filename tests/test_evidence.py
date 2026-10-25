"""Evidence is exact, scoped, portable, and independent of unrelated dirty files."""
from __future__ import annotations

import json
import zipfile

import pytest
from docx import Document

from wongo import evidence
from wongo.engine import render_project
from wongo.errors import InputError
from wongo.review_export import artifact_identity, compare_artifacts, export_review, verify_bundle


@pytest.fixture
def rendered(stub_quarto, wongo_project, monkeypatch):
    monkeypatch.setattr(evidence, "engine_identity", lambda: {"version": "test", "sha256": "engine"})
    render_project(wongo_project, "submission")
    return wongo_project


def _record(project):
    return evidence.make_record(evidence.snapshot(project), "audit", {"id": "context"},
                                {"ok": True, "assertions": []},
                                {"checked": ["one paragraph"], "untested": ["other prose"]})


def test_record_detects_tampering_changed_outputs_and_engine(rendered, monkeypatch):
    record = _record(rendered)
    evidence.save_record(rendered, record)
    assert evidence.verify_record(rendered, record)["state"] == "fresh"
    assert evidence.load_records(rendered, "audit")[0] == record
    modified = json.loads(json.dumps(record))
    modified["result"]["ok"] = False
    assert evidence.verify_record(rendered, modified)["state"] == "invalid"
    monkeypatch.setattr(evidence, "engine_identity", lambda: {"sha256": "different source"})
    assert evidence.verify_record(rendered, record)["state"] == "stale"
    path = rendered / "output" / "main-submission.docx"
    path.write_bytes(path.read_bytes() + b"changed")
    assert evidence.verify_record(rendered, record)["state"] == "stale"


def test_dirty_install_state_is_information_not_identity(rendered, monkeypatch):
    import wongo.engine as engine

    monkeypatch.setattr(engine, "install_identity", lambda: {"state": "dirty", "commit": "same"})
    before = evidence.snapshot(rendered)
    monkeypatch.setattr(engine, "install_identity", lambda: {"state": "clean", "commit": "same"})
    after = evidence.snapshot(rendered)
    assert before["id"] == after["id"]
    assert before["installation"] != after["installation"]


def test_engine_source_change_is_distinguishable_with_the_same_version(tmp_path, monkeypatch):
    package = tmp_path / "package"
    package.mkdir()
    main = package / "evidence.py"
    main.write_text("value = 1\n", encoding="utf-8")
    monkeypatch.setattr(evidence, "__file__", str(main))
    before = evidence.engine_identity()
    main.write_text("value = 2\n", encoding="utf-8")
    after = evidence.engine_identity()
    assert before["version"] == after["version"]
    assert before["sha256"] != after["sha256"]


@pytest.mark.parametrize("raw", ["../secret", "/secret", "C:\\secret", "", ".", "a/../../secret"])
def test_record_artifacts_cannot_escape_their_project(tmp_path, raw):
    with pytest.raises(InputError):
        evidence.safe_path(tmp_path, raw, must_exist=False)


def test_artifact_changes_invalidate_evidence(rendered):
    artifact = rendered / "output" / ".wongo-evidence" / "notes.txt"
    artifact.parent.mkdir(parents=True)
    artifact.write_text("inspected", encoding="utf-8")
    record = evidence.make_record(evidence.snapshot(rendered), "inspection", {}, {"ok": True}, {},
                                  artifacts=[{"path": artifact.relative_to(rendered).as_posix(),
                                              "sha256": evidence.hash_file(artifact)}])
    assert evidence.verify_record(rendered, record)["state"] == "fresh"
    artifact.write_text("changed", encoding="utf-8")
    assert evidence.verify_record(rendered, record)["state"] == "stale"


def test_export_is_allowlisted_and_preserves_coverage(rendered, tmp_path):
    (rendered / ".env").write_text("private-secret-do-not-export", encoding="utf-8")
    (rendered / "raw-data.csv").write_text("private-scientific-data", encoding="utf-8")
    record = _record(rendered)
    evidence.save_record(rendered, record)
    destination = tmp_path / "review.zip"
    export_review(rendered, destination, include_outputs=True)
    assert verify_bundle(destination)["ok"]
    with zipfile.ZipFile(destination) as package:
        assert set(package.namelist()) == {
            "review.json", "files/output/main-submission.docx", "files/output/si-submission.docx",
        }
        manifest = json.loads(package.read("review.json"))
    assert manifest["records"][0]["coverage"]["untested"] == ["other prose"]
    assert manifest["source_files_included"] == []
    assert "private-secret" not in json.dumps(manifest)


def test_old_receipts_cannot_be_attached_to_new_build(rendered, tmp_path):
    record = _record(rendered)
    evidence.save_record(rendered, record)
    source = rendered / "index.qmd"
    source.write_text(source.read_text(encoding="utf-8") + "\nNew material.\n", encoding="utf-8")
    render_project(rendered, "submission")
    with pytest.raises(InputError, match="stale"):
        export_review(rendered, tmp_path / "bad.zip", record_ids=[record["record_id"]])
    result = export_review(rendered, tmp_path / "current.zip")
    assert result["records"] == 0
    assert result["excluded_records"][0]["record_id"] == record["record_id"]


def test_portable_bundle_detects_changed_table_content(rendered, tmp_path):
    destination = tmp_path / "review.zip"
    export_review(rendered, destination, include_outputs=True)
    with zipfile.ZipFile(destination) as package:
        contents = {name: package.read(name) for name in package.namelist()}
    contents["files/output/main-submission.docx"] += b"table tampered"
    with zipfile.ZipFile(destination, "w") as package:
        for name, raw in contents.items():
            package.writestr(name, raw)
    with pytest.raises(InputError, match="identity mismatch"):
        verify_bundle(destination)


def _docx_rewrite(path, change):
    with zipfile.ZipFile(path) as package:
        parts = {name: package.read(name) for name in package.namelist()}
    change(parts)
    with zipfile.ZipFile(path, "w") as package:
        for name, raw in parts.items():
            package.writestr(name, raw)


def test_docx_date_normalization_keeps_cells_fields_and_alt_text(tmp_path):
    first, second = tmp_path / "first.docx", tmp_path / "second.docx"
    doc = Document()
    doc.add_table(rows=1, cols=1).cell(0, 0).text = "Scientific cell 42"
    doc.save(first)
    second.write_bytes(first.read_bytes())
    _docx_rewrite(second, lambda parts: parts.update({"docProps/core.xml": parts["docProps/core.xml"].replace(
        b"2013-12-23T23:15:00Z", b"2026-10-23T23:15:00Z")}))
    assert compare_artifacts(first, second)["classification"] == "metadata-only"
    assert artifact_identity(first)["sha256"] != artifact_identity(second)["sha256"]
    _docx_rewrite(second, lambda parts: parts.update({"word/document.xml": parts["word/document.xml"].replace(
        b"Scientific cell 42", b"Scientific cell 43")}))
    assert compare_artifacts(first, second)["classification"] == "changed"


def test_pdf_normalization_changes_only_uncompressed_info_dates(tmp_path):
    first, second = tmp_path / "first.pdf", tmp_path / "second.pdf"
    raw = _pdf(b"/CreationDate (D:20261022000000Z) /ModDate (D:20261022000000Z)")
    first.write_bytes(raw)
    second.write_bytes(raw.replace(b"D:20261022", b"D:20261023"))
    assert compare_artifacts(first, second)["classification"] == "metadata-only"
    second.write_bytes(second.read_bytes().replace(b"Scientific 42", b"Scientific 43"))
    assert compare_artifacts(first, second)["classification"] == "changed"


def test_pdf_date_words_in_titles_and_streams_are_not_erased(tmp_path):
    first, second = tmp_path / "first.pdf", tmp_path / "second.pdf"
    raw = (b"%PDF-1.4\n1 0 obj\n<< /Length 42 >>\nstream\n/CreationDate (D:20261022000000Z)\nendstream\nendobj\n"
           b"2 0 obj\n<< /Title (/CreationDate \\(D:20261022000000Z\\)) >>\nendobj\n"
           b"trailer\n<< /Root 1 0 R /Info 2 0 R >>\n%%EOF\n")
    first.write_bytes(raw)
    second.write_bytes(raw.replace(b"D:20261022", b"D:20261023"))
    assert compare_artifacts(first, second)["classification"] == "changed"


def _pdf(info):
    raw = b"%PDF-1.4\n"
    first = len(raw)
    raw += b"1 0 obj\n<< /Length 14 >>\nstream\nScientific 42\nendstream\nendobj\n"
    second = len(raw)
    raw += b"2 0 obj\n<< " + info + b" >>\nendobj\n"
    xref = len(raw)
    raw += (f"xref\n0 3\n0000000000 65535 f \n{first:010} 00000 n \n"
            f"{second:010} 00000 n \ntrailer\n<< /Size 3 /Root 1 0 R /Info 2 0 R >>\n"
            f"startxref\n{xref}\n%%EOF\n").encode()
    return raw


def test_pdf_info_looking_stream_content_has_no_metadata_privilege(tmp_path):
    first, second = tmp_path / "first.pdf", tmp_path / "second.pdf"
    raw = _pdf(b"/Title (Scientific result)")
    raw = raw.replace(b"Scientific 42", b"/CreationDate (D:20261022000000Z)")
    first.write_bytes(raw)
    second.write_bytes(raw.replace(b"D:20261022", b"D:20261023"))
    assert compare_artifacts(first, second)["classification"] == "changed"


@pytest.mark.parametrize("old,new", [(b'PAGE', b'NUMPAGES'), (b'Figure description', b'Incorrect description')])
def test_docx_field_values_and_accessibility_are_content(tmp_path, old, new):
    first, second = tmp_path / "first.docx", tmp_path / "second.docx"
    doc = Document()
    doc.add_paragraph("Scientific result")
    doc.save(first)
    def add_fields(parts):
        parts["word/document.xml"] = parts["word/document.xml"].replace(
            b"</w:body>", b'<w:p><w:fldSimple w:instr="PAGE"/></w:p>'
            b'<w:p><w:r><w:drawing><wp:docPr xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" '
            b'id="1" name="Figure" descr="Figure description"/></w:drawing></w:r></w:p></w:body>')
    _docx_rewrite(first, add_fields)
    second.write_bytes(first.read_bytes())
    _docx_rewrite(second, lambda parts: parts.update({"word/document.xml": parts["word/document.xml"].replace(old, new)}))
    assert compare_artifacts(first, second)["classification"] == "changed"


def test_export_carries_history_as_provenance_only(rendered, tmp_path):
    old = _record(rendered)
    evidence.save_record(rendered, old)
    source = rendered / "index.qmd"
    source.write_text(source.read_text(encoding="utf-8") + "\nNew draft.\n", encoding="utf-8")
    render_project(rendered, "submission")
    current = evidence.make_record(evidence.snapshot(rendered), "visual-review", {},
                                    {"carried_from": old["record_id"]}, {})
    evidence.save_record(rendered, current)
    destination = tmp_path / "history.zip"
    result = export_review(rendered, destination, record_ids=[current["record_id"]])
    assert result["records"] == 1 and result["historical_records"] == 1
    assert verify_bundle(destination)["historical_records"] == 1
    with zipfile.ZipFile(destination) as package:
        manifest = json.loads(package.read("review.json"))
    assert manifest["records"] == [current]
    assert manifest["historical_records"] == [old]


def test_export_refuses_missing_reference_and_reports_omitted_artifacts(rendered, tmp_path):
    record = evidence.make_record(evidence.snapshot(rendered), "gate-proposal", {},
                                    {"evidence_ids": ["missing"]}, {})
    evidence.save_record(rendered, record)
    with pytest.raises(InputError, match="missing record"):
        export_review(rendered, tmp_path / "broken.zip")
    path = evidence.evidence_dir(rendered) / "gate-proposal" / f"{record['record_id']}.json"
    path.unlink()
    artifact = rendered / "raw-data.csv"
    artifact.write_text("private", encoding="utf-8")
    note = evidence.make_record(evidence.snapshot(rendered), "review-note", {}, {}, {},
                                 artifacts=[{"path": "raw-data.csv", "sha256": evidence.hash_file(artifact)}])
    evidence.save_record(rendered, note)
    result = export_review(rendered, tmp_path / "omitted.zip")
    assert result["omitted_artifacts"] == ["raw-data.csv"]
    assert verify_bundle(tmp_path / "omitted.zip")["omitted_artifacts"] == ["raw-data.csv"]


def test_rehashed_bundle_cannot_attach_changed_output_to_old_subject(rendered, tmp_path):
    import hashlib

    path = tmp_path / "review.zip"
    export_review(rendered, path, include_outputs=True)
    with zipfile.ZipFile(path) as package:
        parts = {name: package.read(name) for name in package.namelist()}
    manifest = json.loads(parts["review.json"])
    item = manifest["files"][0]
    changed = parts[item["path"]] + b"changed scientific text"
    parts[item["path"]] = changed
    item.update(sha256=hashlib.sha256(changed).hexdigest(), bytes=len(changed))
    item["identity"]["sha256"] = item["sha256"]
    manifest["bundle_id"] = evidence.digest({k: v for k, v in manifest.items() if k != "bundle_id"})
    parts["review.json"] = json.dumps(manifest).encode()
    with zipfile.ZipFile(path, "w") as package:
        for name, raw in parts.items():
            package.writestr(name, raw)
    with pytest.raises(InputError, match="sealed build"):
        verify_bundle(path)


def test_snapshot_refuses_source_save_after_initial_freshness_check(rendered, monkeypatch):
    import wongo.engine as engine

    original = engine.output_freshness
    def changed(project, target):
        result = original(project, target)
        source = project / "index.qmd"
        source.write_text(source.read_text(encoding="utf-8") + "\nConcurrent save.\n", encoding="utf-8")
        return result
    monkeypatch.setattr(engine, "output_freshness", changed)
    with pytest.raises(InputError, match="changed while capturing evidence"):
        evidence.snapshot(rendered)


@pytest.mark.parametrize("location", ["output/.wongo-evidence", "output/.wongo-evidence/audit"])
def test_evidence_storage_refuses_external_symlink(rendered, tmp_path, location):
    record = _record(rendered)
    outside = tmp_path / "outside"
    outside.mkdir()
    link = rendered / location
    link.parent.mkdir(parents=True, exist_ok=True)
    link.symlink_to(outside, target_is_directory=True)
    with pytest.raises(InputError, match="escapes"):
        evidence.save_record(rendered, record)
    assert list(outside.iterdir()) == []
    with pytest.raises(InputError, match="escapes"):
        if location.endswith("/audit"):
            # Store traversal must also reject an externally linked child when reading.
            (outside / f"{record['record_id']}.json").write_text(json.dumps(record), encoding="utf-8")
        evidence.load_records(rendered)


def test_malformed_record_kind_is_actionable():
    with pytest.raises(InputError, match="Unknown evidence kind"):
        evidence.validate_record({"schema": evidence.SCHEMA, "kind": []})


def test_malformed_portable_manifest_is_actionable(tmp_path):
    path = tmp_path / "invalid.zip"
    with zipfile.ZipFile(path, "w") as package:
        package.writestr("review.json", "[]")
    with pytest.raises(InputError, match="mapping"):
        verify_bundle(path)


def test_malformed_page_references_are_actionable():
    from wongo.review_export import record_references

    with pytest.raises(InputError, match="page mappings"):
        record_references({"result": {"checked_pages": [1]}})
