"""Context-aware helpers report what was checked without deciding scientific policy."""
from __future__ import annotations

import hashlib
import struct
import zlib

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from wongo.audit_helpers import (
    figure_order_assertion,
    paragraph_assertion,
    scalar_assertion,
    table_cell_assertion,
)


@pytest.mark.parametrize("data,pointer,allow_missing,ok", [
    ({}, "/value", True, False),
    ({"value": None}, "/value", False, False),
    ({"value": None}, "/value", True, True),
    ({"value": ""}, "/value", True, False),
    ({"value": []}, "/value", False, False),
    ({"value": [1]}, "/value", False, False),
    ({"value": float("nan")}, "/value", False, False),
    ({"value": float("inf")}, "/value", False, False),
    ({"value": 12.4}, "/value", False, True),
    ({"value": "NA"}, "/value", False, True),
    ({"values": [{"a/b": 2}]}, "/values/0/a~1b", False, True),
    ({"values": [1]}, "/values/-1", False, False),
    ({"values": [1]}, "/values/00", False, False),
])
def test_scalar_requirements_are_explicit(data, pointer, allow_missing, ok):
    result = scalar_assertion(data, pointer, allow_missing=allow_missing)
    assert result["ok"] is ok
    assert result["artifact"] == pointer
    assert result["detail"]


def test_value_in_another_sentence_does_not_satisfy_the_expected_context(tmp_path):
    doc = Document()
    doc.add_paragraph("Observed current was 13.1 A.")
    doc.add_paragraph("A different treatment yielded 12.4 A.")
    path = tmp_path / "main.docx"
    doc.save(path)
    failure = paragraph_assertion(path, "Observed current", "12.4")
    assert not failure["ok"]
    assert failure["rendered_context"]["text"] == "Observed current was 13.1 A."
    assert paragraph_assertion(path, "Observed current", "13.1")["ok"]


def test_another_sentence_in_the_same_paragraph_cannot_supply_the_value(tmp_path):
    doc = Document()
    doc.add_paragraph("Yield was unavailable. Recovery was 42 %.")
    path = tmp_path / "main.docx"
    doc.save(path)
    result = paragraph_assertion(path, "Yield was", "42")
    assert not result["ok"]
    assert result["actual"] == "Yield was unavailable."
    assert result["rendered_context"]["text"] == "Yield was unavailable."
    assert result["rendered_context"]["paragraph_text"] == "Yield was unavailable. Recovery was 42 %."
    assert paragraph_assertion(path, "Recovery was", "42")["ok"]


@pytest.mark.parametrize("actual,expected,ok", [
    ("-12.4", "12.4", False),
    ("+12.4", "12.4", False),
    ("−12.4", "12.4", False),
    ("-12.4", "-12.4", True),
    ("12.4e3", "12.4", False),
    ("12.4e-3", "12.4", False),
    ("12.4E+3", "12.4", False),
    ("12.4e-3", "3", False),
    ("12.4E+3", "3", False),
    ("12.4e-3", "12.4e-3", True),
    ("12.4", "12.4", True),
])
def test_signed_numeric_and_exponent_tokens_are_exact(tmp_path, actual, expected, ok):
    doc = Document()
    doc.add_paragraph(f"Observed current was {actual} A.")
    path = tmp_path / "main.docx"
    doc.save(path)
    assert paragraph_assertion(path, "Observed current", expected)["ok"] is ok


def test_full_expected_sentence_matches_its_punctuation_span(tmp_path):
    doc = Document()
    doc.add_paragraph("Before this sentence. Observed current was 12.4 A. After this sentence.")
    path = tmp_path / "main.docx"
    doc.save(path)
    result = paragraph_assertion(path, "Observed current", "Observed current was 12.4 A.")
    assert result["ok"]
    assert result["actual"] == "Observed current was 12.4 A."


def test_repeated_anchor_within_one_paragraph_is_ambiguous(tmp_path):
    doc = Document()
    doc.add_paragraph("Current was unavailable. Current was 12.4 A.")
    path = tmp_path / "main.docx"
    doc.save(path)
    assert not paragraph_assertion(path, "Current was", "12.4")["ok"]


def test_context_must_be_unique_and_numeric_tokens_are_exact(tmp_path):
    doc = Document()
    doc.add_paragraph("Current was 112 A.")
    path = tmp_path / "main.docx"
    doc.save(path)
    assert not paragraph_assertion(path, "Current", "12")["ok"]
    doc.add_paragraph("Current was 12 A.")
    doc.save(path)
    assert not paragraph_assertion(path, "Current", "12")["ok"]


def test_value_at_sentence_end_is_not_confused_with_decimal_prefix(tmp_path):
    doc = Document()
    doc.add_paragraph("Observed current was 12.4.")
    path = tmp_path / "main.docx"
    doc.save(path)
    assert paragraph_assertion(path, "Observed current", "12.4")["ok"]
    assert not paragraph_assertion(path, "Observed current", "12")["ok"]


def test_deleted_word_value_does_not_count_as_visible_content(tmp_path):
    doc = Document()
    paragraph = doc.add_paragraph("Observed current was 13.1 A.")
    deletion = OxmlElement("w:del")
    deletion.set(qn("w:id"), "1")
    run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = "12.4"
    run.append(text)
    deletion.append(run)
    paragraph._p.append(deletion)
    path = tmp_path / "main.docx"
    doc.save(path)
    assert not paragraph_assertion(path, "Observed current", "12.4")["ok"]


def test_table_values_are_checked_in_the_requested_cell(tmp_path):
    doc = Document()
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Treatment"
    table.cell(0, 1).text = "Current"
    table.cell(1, 0).text = "A"
    table.cell(1, 1).text = "12.4"
    path = tmp_path / "main.docx"
    doc.save(path)
    assert table_cell_assertion(path, 0, 1, 1, "12.4")["ok"]
    assert not table_cell_assertion(path, 0, 1, 0, "12.4")["ok"]
    assert not table_cell_assertion(path, 0, 5, 0, "12.4")["ok"]


def _png(path, color):
    def chunk(tag, data):
        payload = tag + data
        return struct.pack(">I", len(data)) + payload + struct.pack(">I", zlib.crc32(payload) & 0xFFFFFFFF)

    raw = (b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", struct.pack(">IIBBBBB", 1, 1, 8, 2, 0, 0, 0))
           + chunk(b"IDAT", zlib.compress(b"\x00" + bytes(color))) + chunk(b"IEND", b""))
    path.write_bytes(raw)
    return hashlib.sha256(raw).hexdigest()


def test_embedded_figure_identity_and_order_are_exact(tmp_path):
    doc = Document()
    first = tmp_path / "first.png"
    second = tmp_path / "second.png"
    first_hash = _png(first, (255, 0, 0))
    second_hash = _png(second, (0, 0, 255))
    doc.add_picture(str(first))
    doc.add_picture(str(second))
    path = tmp_path / "main.docx"
    doc.save(path)
    assert figure_order_assertion(path, [first_hash, second_hash])["ok"]
    assert not figure_order_assertion(path, [second_hash, first_hash])["ok"]


def test_deleted_image_is_not_part_of_visible_figure_order(tmp_path):
    doc = Document()
    image = tmp_path / "deleted.png"
    _png(image, (255, 0, 0))
    run = doc.add_paragraph().add_run()
    run.add_picture(str(image))
    deletion = OxmlElement("w:del")
    deletion.set(qn("w:id"), "1")
    run._r.addprevious(deletion)
    deletion.append(run._r)
    path = tmp_path / "main.docx"
    doc.save(path)
    assert figure_order_assertion(path, [])["ok"]
