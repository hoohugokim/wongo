"""Declared inputs, document-scoped cross-references, ES&T rule coverage,
and the in-transaction postprocess stage. Synthetic projects only."""
from __future__ import annotations

import importlib.util
import json
import os
import zipfile
from datetime import date
from pathlib import Path

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from wongo.engine import output_freshness, postprocess_main, render_project, source_fingerprint
from wongo.engine.checks import rule_coverage, run_checks, word_count
from wongo.errors import GateError, InputError
from wongo.profiles import load_profile, manuscript_type
from wongo.profiles.verify import si_placement_problems, verify_profile

_helpers_spec = importlib.util.spec_from_file_location(
    "wongo_test_run_checks_helpers",
    Path(__file__).with_name("test_run_checks.py"),
)
assert _helpers_spec is not None and _helpers_spec.loader is not None
_helpers = importlib.util.module_from_spec(_helpers_spec)
_helpers_spec.loader.exec_module(_helpers)
_check = _helpers._check
make_project = _helpers.make_project


def _journal(project: Path, text: str) -> None:
    (project / "_journal.yml").write_text(text, encoding="utf-8")


def test_declared_json_and_helper_changes_go_stale_and_unrelated_files_do_not(
        stub_quarto, wongo_project, monkeypatch):
    data = wongo_project / "data"
    scripts = wongo_project / "scripts"
    data.mkdir()
    scripts.mkdir()
    (data / "facts.json").write_text('{"n": 1}\n', encoding="utf-8")
    (scripts / "setup.R").write_text("helper <- function() 1\n", encoding="utf-8")
    _journal(
        wongo_project,
        "journal: demo\nms_type: article\nstyle: default\n"
        "dependencies:\n  - data/facts.json\n  - scripts/setup.R\n",
    )
    render_project(wongo_project, "collab")
    assert output_freshness(wongo_project, "collab")["state"] == "fresh"

    (data / "facts.json").write_text('{"n": 2}\n', encoding="utf-8")
    assert "data/facts.json" in output_freshness(wongo_project, "collab")["changed"]
    (data / "facts.json").write_text('{"n": 1}\n', encoding="utf-8")
    (scripts / "setup.R").write_text("helper <- function() 2\n", encoding="utf-8")
    assert "scripts/setup.R" in output_freshness(wongo_project, "collab")["changed"]

    (scripts / "setup.R").write_text("helper <- function() 1\n", encoding="utf-8")
    before = source_fingerprint(wongo_project)
    (wongo_project / "notes.txt").write_text("unrelated\n", encoding="utf-8")
    assert source_fingerprint(wongo_project) == before
    assert output_freshness(wongo_project, "collab")["state"] == "fresh"


def test_a_missing_declared_input_names_its_path(wongo_project):
    _journal(
        wongo_project,
        "journal: demo\nms_type: article\nstyle: default\n"
        "dependencies:\n  - data/facts.json\n",
    )
    check = _check(run_checks(wongo_project), "dependencies")
    assert check.level == "HARD" and not check.ok
    assert "data/facts.json" in check.detail
    with pytest.raises(InputError, match="data/facts.json"):
        render_project(wongo_project, "collab")


def test_an_external_dependency_is_hashed_and_not_copied(stub_quarto, wongo_project, tmp_path):
    outside = tmp_path / "outside.json"
    outside.write_text('{"n": 1}\n', encoding="utf-8")
    rel = Path(os.path.relpath(outside, wongo_project)).as_posix()
    _journal(
        wongo_project,
        "journal: demo\nms_type: article\nstyle: default\n"
        "dependencies:\n"
        f"  - path: {rel}\n    external: true\n",
    )
    render_project(wongo_project, "collab")
    assert not (wongo_project / "outside.json").exists()
    fingerprint = source_fingerprint(wongo_project)
    assert outside.resolve().as_posix() in fingerprint
    outside.write_text('{"n": 2}\n', encoding="utf-8")
    assert output_freshness(wongo_project, "collab")["state"] == "stale"


def test_discovery_reports_a_literal_read_without_adding_it_to_the_fingerprint(
        stub_quarto, wongo_project):
    (wongo_project / "index.qmd").write_text(
        "---\ntitle: T\nabstract: Short abstract here.\n---\n\n"
        "```{r}\nread_csv('data/facts.json')\n```\n",
        encoding="utf-8",
    )
    _journal(
        wongo_project,
        "journal: demo\nms_type: article\nstyle: default\n"
        "dependencies_discover: true\n",
    )
    check = _check(run_checks(wongo_project), "undeclared-inputs")
    assert not check.ok
    assert "data/facts.json" in check.detail
    assert "not exhaustive" in check.detail
    before = source_fingerprint(wongo_project)
    (wongo_project / "data").mkdir()
    (wongo_project / "data" / "facts.json").write_text("{}\n", encoding="utf-8")
    assert source_fingerprint(wongo_project) == before


def test_a_manifest_without_profile_identity_is_stale(stub_quarto, wongo_project):
    render_project(wongo_project, "collab")
    path = wongo_project / "output" / ".wongo-manifest.json"
    manifest = json.loads(path.read_text(encoding="utf-8"))
    del manifest["targets"]["collab"]["identity"]
    path.write_text(json.dumps(manifest), encoding="utf-8")
    fresh = output_freshness(wongo_project, "collab")
    assert fresh["state"] == "stale"
    assert "profile identity" in fresh["changed"]
    assert "does not name profile or style identity" in fresh["detail"]


def test_main_reference_defined_only_in_si_fails(tmp_path):
    project = make_project(
        tmp_path,
        index_body="See @tbl-si-only in the supplement.",
        include_si=True,
        si_body="Only here.\n\n{#tbl-si-only}\n",
    )
    check = _check(run_checks(project), "crossrefs")
    assert not check.ok
    assert "index.qmd: @tbl-si-only" in check.detail


def test_a_local_reference_and_an_si_internal_reference_pass(tmp_path):
    project = make_project(
        tmp_path,
        index_body="See @tbl-main.\n\n{#tbl-main}\n",
        include_si=True,
        si_body="See @tbl-extra.\n\n{#tbl-extra}\n",
    )
    assert _check(run_checks(project), "crossrefs").ok
    assert _check(run_checks(project), "uncited").ok


def test_the_same_label_in_both_documents_fails(tmp_path):
    project = make_project(
        tmp_path,
        index_body="{#fig-shared}\n",
        include_si=True,
        si_body="{#fig-shared}\n",
    )
    check = _check(run_checks(project), "crossrefs")
    assert not check.ok
    assert "fig-shared" in check.detail


def test_a_cross_reference_inside_a_code_fence_is_not_a_callout(tmp_path):
    project = make_project(
        tmp_path,
        index_body="```\nSee @fig-example in this snippet.\n```\n",
        include_si=False,
    )
    assert _check(run_checks(project), "crossrefs").ok


def test_quarto_unresolved_crossref_fails_submission_and_is_kept_on_collab(
        stub_quarto, wongo_project, monkeypatch):
    monkeypatch.setenv(
        "WONGO_STUB_QUARTO_STDOUT",
        "WARNING: Unable to resolve crossref @tbl-si-only\n",
    )
    with pytest.raises(GateError) as caught:
        render_project(wongo_project, "submission")
    assert any(c.name == "crossrefs-rendered" and not c.ok for c in caught.value.details["checks"])
    assert not (wongo_project / "output" / "main-submission.docx").exists()

    result = render_project(wongo_project, "collab")
    assert (wongo_project / "output" / "main-collab.docx").is_file()
    assert any(c.name == "crossrefs-rendered" and not c.ok for c in result.checks)
    manifest = json.loads((wongo_project / "output" / ".wongo-manifest.json").read_text(encoding="utf-8"))
    assert "crossrefs-rendered" in manifest["targets"]["collab"]["checks_failed"]


def _document_xml(path: Path) -> bytes:
    with zipfile.ZipFile(path) as package:
        return package.read("word/document.xml")


def _est_project(tmp_path: Path, abstract: str, body: str, ms_type: str = "research-article") -> Path:
    project = tmp_path / "est"
    project.mkdir(parents=True)
    (project / "_journal.yml").write_text(
        f"journal: est\nms_type: {ms_type}\nstyle: default\n", encoding="utf-8")
    (project / "_quarto.yml").write_text("bibliography: refs.bib\n", encoding="utf-8")
    (project / "refs.bib").write_text(
        "@article{doe2020, title = {T}, author = {Doe}, year = {2020}}\n", encoding="utf-8")
    (project / "index.qmd").write_text(
        f"---\ntitle: T\nabstract: {abstract}\n---\n\n{body}\n", encoding="utf-8")
    return project


def test_est_abstract_bounds(tmp_path):
    short = _check(run_checks(_est_project(tmp_path, "two words", "Body.")), "abstract-word-limit")
    assert not short.ok and "below" in short.detail
    long = _check(run_checks(_est_project(
        tmp_path / "long", " ".join(f"w{i}" for i in range(250)), "Body.")), "abstract-word-limit")
    assert not long.ok and "exceed" in long.detail
    for count, folder in ((150, "low"), (200, "high")):
        words = " ".join(f"w{i}" for i in range(count))
        check = _check(run_checks(_est_project(tmp_path / folder, words, "Body.")), "abstract-word-limit")
        assert check.ok, check.detail
    missing = tmp_path / "missing"
    project = _est_project(missing, "placeholder", "Body.")
    text = (project / "index.qmd").read_text(encoding="utf-8").replace("abstract: placeholder\n", "")
    (project / "index.qmd").write_text(text, encoding="utf-8")
    assert not _check(run_checks(project), "abstract-word-limit").ok


def test_est_acknowledgment_is_excluded_and_a_caption_word_counts():
    profile = load_profile("est")
    policy = manuscript_type(profile, "research-article")["word_count"]
    text = (
        "---\nabstract: alpha beta\n---\n\n# Introduction\n\nOne two three.\n\n"
        "![capword](figures/demo.png)\n\n# Acknowledgment\n\nThanks to Ada.\n"
    )
    counted = word_count(text, **policy)
    assert counted == word_count(text) - 3
    without_caption = text.replace("![capword](figures/demo.png)\n", "")
    assert word_count(without_caption, **policy) == counted - 1


def test_est_feature_abstract_length_stays_manual():
    profile = load_profile("est")
    rules = {
        item["id"]: item["enforcement"]
        for item in rule_coverage(profile, manuscript_type(profile, "feature"))
    }
    assert rules["abstract-length"] == "manual"
    assert rules["word-count-exclusions"] == "machine"
    research = {
        item["id"]: item["enforcement"]
        for item in rule_coverage(profile, manuscript_type(profile, "research-article"))
    }
    assert research["abstract-length"] == "machine"
    assert research["abstract-one-paragraph"] == "manual"
    assert research["keywords"] == "manual"


def test_est_skill_and_checklist_match_the_profile_placement():
    profile = load_profile("est")
    assert si_placement_problems(Path(profile["_dir"]), profile) == []


def test_a_checklist_that_contradicts_si_placement_fails_verify(tmp_path, monkeypatch):
    from wongo import profiles

    root = tmp_path / "profiles"
    (root / "demo" / "references").mkdir(parents=True)
    today = date.today().isoformat()
    (root / "demo" / "profile.yml").write_text(
        "slug: demo\njournal: Demo\n"
        "manuscript_types: [{type: article, word_limit: 100, counting_rule: r}]\n"
        "si: {availability_paragraph: before-acknowledgments}\n"
        f"sources: [https://example.org]\nverified_date: {today}\n",
        encoding="utf-8",
    )
    (root / "demo" / "references" / "submission-checklist.md").write_text(
        "The availability paragraph goes after Acknowledgments.\n", encoding="utf-8")
    monkeypatch.setattr(profiles, "candidate_dirs", lambda project=None: [root])
    result = verify_profile("demo", offline=True, emit=lambda *_args: None)
    assert not result.ok
    assert any("after-acknowledgments" in problem for problem in result.problems)


def _png(path: Path) -> Path:
    import struct
    import zlib

    def chunk(tag: bytes, data: bytes) -> bytes:
        body = tag + data
        return struct.pack(">I", len(data)) + body + struct.pack(">I", zlib.crc32(body) & 0xFFFFFFFF)

    raw = b"\x00" + b"\x40\x80\xc0"
    path.write_bytes(
        b"\x89PNG\r\n\x1a\n"
        + chunk(b"IHDR", struct.pack(">IIBBBBB", 1, 1, 8, 2, 0, 0, 0))
        + chunk(b"IDAT", zlib.compress(raw))
        + chunk(b"IEND", b"")
    )
    return path


def test_engine_postprocess_keeps_alt_text_revisions_and_math(wongo_project, tmp_path):
    from wongo.profiles import load_journal_config, load_profile

    image = _png(tmp_path / "pixel.png")
    doc = Document()
    doc.add_paragraph("Body")
    run = doc.add_paragraph().add_run()
    run.add_picture(str(image))
    for element in doc.element.iter():
        if element.tag.endswith("docPr"):
            element.set("descr", "A reactor schematic")
    inserted = OxmlElement("w:ins")
    inserted.set(qn("w:id"), "1")
    inserted.set(qn("w:author"), "Ada")
    hold = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = "kept"
    hold.append(text)
    inserted.append(hold)
    doc.paragraphs[0]._p.append(inserted)
    math = OxmlElement("m:oMath")
    doc.paragraphs[0]._p.append(math)
    path = wongo_project / "output" / "main-collab.docx"
    path.parent.mkdir(parents=True)
    doc.save(path)
    cfg = load_journal_config(wongo_project)
    profile = load_profile(cfg["journal"], wongo_project)
    postprocess_main(path, profile, cfg, "collab", wongo_project)
    xml = _document_xml(path)
    assert b"A reactor schematic" in xml
    assert b"w:ins" in xml or b"<w:ins" in xml
    assert b"oMath" in xml


def _script(project: Path, name: str, body: str) -> None:
    path = project / "scripts" / name
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(body, encoding="utf-8")


def test_a_registered_postprocess_is_fresh_and_a_failed_one_keeps_the_previous_output(
        stub_quarto, wongo_project):
    render_project(wongo_project, "collab")
    previous = (wongo_project / "output" / "main-collab.docx").read_bytes()
    _script(wongo_project, "fail.py", "import sys\nsys.exit(1)\n")
    _journal(
        wongo_project,
        "journal: demo\nms_type: article\nstyle: default\n"
        "postprocess:\n  - scripts/fail.py\n",
    )
    with pytest.raises(InputError, match="fail.py"):
        render_project(wongo_project, "collab")
    assert (wongo_project / "output" / "main-collab.docx").read_bytes() == previous

    _script(
        wongo_project, "mark.py",
        "import os\nfrom docx import Document\n"
        "for raw in os.environ['WONGO_POSTPROCESS_OUTPUTS'].split(os.pathsep):\n"
        "    doc = Document(raw)\n"
        "    doc.add_paragraph('WONGO_MARK')\n"
        "    doc.save(raw)\n",
    )
    _journal(
        wongo_project,
        "journal: demo\nms_type: article\nstyle: default\n"
        "postprocess:\n  - scripts/mark.py\n",
    )
    render_project(wongo_project, "collab")
    produced = wongo_project / "output" / "main-collab.docx"
    assert b"WONGO_MARK" in _document_xml(produced)
    fresh = output_freshness(wongo_project, "collab")
    assert fresh["state"] == "fresh"
    manifest = json.loads((wongo_project / "output" / ".wongo-manifest.json").read_text(encoding="utf-8"))
    entry = manifest["targets"]["collab"]
    assert entry["postprocess"][0]["path"] == "scripts/mark.py"
    assert entry["postprocess"][0]["outputs"]["main-collab.docx"]["before"] != (
        entry["postprocess"][0]["outputs"]["main-collab.docx"]["after"]
    )

    produced.write_bytes(produced.read_bytes() + b" ")
    assert output_freshness(wongo_project, "collab")["state"] == "modified"
    render_project(wongo_project, "collab")
    assert output_freshness(wongo_project, "collab")["state"] == "fresh"
    (wongo_project / "scripts" / "mark.py").write_text("# changed\n", encoding="utf-8")
    assert output_freshness(wongo_project, "collab")["state"] == "stale"
