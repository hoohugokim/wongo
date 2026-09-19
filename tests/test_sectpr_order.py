"""Regression tests recreated for the sectPr child-order fix in
src/wongo/docxpatch/__init__.py (add_line_numbers / restart_page_numbering),
pinning the schema-order bug documented in docs/docx-quirks.md (2026-07-03):
CT_SectPr requires w:lnNumType/w:pgNumType to precede w:cols (and its
ECMA-376 successors); a plain `.append()` placed them after, which Word
treats as corrupt XML. The originals were lost in a migration; see
notes/lost-tests-recreated.md for provenance.
"""
from __future__ import annotations

from docx import Document
from docx.oxml.ns import qn

from wongo.docxpatch import add_line_numbers, restart_page_numbering


def _child_index(sect_pr, tag):
    for i, child in enumerate(sect_pr):
        if child.tag == qn(tag):
            return i
    return None


def test_add_line_numbers_places_lnNumType_before_cols():
    doc = Document()
    sect_pr = doc.sections[0]._sectPr
    assert _child_index(sect_pr, "w:cols") is not None  # blank template ships w:cols

    add_line_numbers(doc)

    ln_idx = _child_index(sect_pr, "w:lnNumType")
    cols_idx = _child_index(sect_pr, "w:cols")
    assert ln_idx is not None
    assert cols_idx is not None
    assert ln_idx < cols_idx


def test_restart_page_numbering_places_pgNumType_before_cols():
    doc = Document()
    sect_pr = doc.sections[0]._sectPr

    restart_page_numbering(doc)

    pg_idx = _child_index(sect_pr, "w:pgNumType")
    cols_idx = _child_index(sect_pr, "w:cols")
    assert pg_idx is not None
    assert cols_idx is not None
    assert pg_idx < cols_idx


def test_restart_then_add_line_numbers_yields_schema_order():
    doc = Document()
    sect_pr = doc.sections[0]._sectPr

    # restart_page_numbering runs FIRST (as wongo's postprocess does), so
    # w:pgNumType already exists in the sectPr by the time add_line_numbers
    # computes lnNumType's insertion point — this is exactly the scenario
    # the docs/docx-quirks.md 2026-07-03 entry calls out as a subtle second
    # bug (lnNumType's successor list must include w:pgNumType, not just the
    # shared _SECT_PR_TAIL).
    restart_page_numbering(doc)
    add_line_numbers(doc)

    ln_idx = _child_index(sect_pr, "w:lnNumType")
    pg_idx = _child_index(sect_pr, "w:pgNumType")
    cols_idx = _child_index(sect_pr, "w:cols")
    assert ln_idx < pg_idx < cols_idx


def test_calling_both_twice_does_not_duplicate_elements():
    doc = Document()
    sect_pr = doc.sections[0]._sectPr

    add_line_numbers(doc)
    restart_page_numbering(doc)
    add_line_numbers(doc)
    restart_page_numbering(doc)

    assert len(sect_pr.findall(qn("w:lnNumType"))) == 1
    assert len(sect_pr.findall(qn("w:pgNumType"))) == 1
