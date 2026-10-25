"""Final DOCX inspection and renderer-specific preview evidence."""
from __future__ import annotations

import hashlib
import os
import struct
import zipfile
import zlib

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from wongo.errors import InputError
from wongo.qa import inspect_docx


def _png(path):
    def chunk(tag, value):
        body = tag + value
        return struct.pack(">I", len(value)) + body + struct.pack(">I", zlib.crc32(body) & 0xFFFFFFFF)
    path.write_bytes(b"\x89PNG\r\n\x1a\n"
                     + chunk(b"IHDR", struct.pack(">IIBBBBB", 1, 1, 8, 2, 0, 0, 0))
                     + chunk(b"IDAT", zlib.compress(b"\x00\x40\x80\xc0"))
                     + chunk(b"IEND", b""))
    return path


def rich_docx(path, *, description="A verified reactor schematic"):
    doc = Document()
    doc.add_paragraph("A long caption " * 35, style="Caption")
    paragraph = doc.add_paragraph("Thin\u2009space and nonbreaking\u00a0space remain.")
    run = paragraph.add_run("2")
    run.font.superscript = True
    inserted = OxmlElement("w:ins")
    inserted.set(qn("w:id"), "1")
    inserted.set(qn("w:author"), "Ada")
    revision_run = OxmlElement("w:r")
    revision_text = OxmlElement("w:t")
    revision_text.text = "A tracked revision."
    revision_run.append(revision_text)
    inserted.append(revision_run)
    paragraph._p.append(inserted)
    math = OxmlElement("m:oMath")
    superscript = OxmlElement("m:sSup")
    for role, value in (("m:e", "x"), ("m:sup", "2")):
        holder = OxmlElement(role)
        math_run = OxmlElement("m:r")
        math_text = OxmlElement("m:t")
        math_text.text = value
        math_run.append(math_text)
        holder.append(math_run)
        superscript.append(holder)
    math.append(superscript)
    paragraph._p.append(math)
    table = doc.add_table(rows=85, cols=9)
    for number, row in enumerate(table.rows, start=1):
        for column, cell in enumerate(row.cells, start=1):
            cell.text = f"r{number}c{column}"
    doc.add_paragraph().add_run().add_picture(str(_png(path.with_suffix(".png"))))
    for element in doc.element.iter():
        if element.tag.endswith("}docPr") and description is not None:
            element.set("descr", description)
    doc.save(path)
    return path


def test_inspection_reads_final_ooxml_and_never_changes_the_docx(tmp_path):
    path = rich_docx(tmp_path / "native.docx")
    before = path.read_bytes()
    result = inspect_docx(path)
    assert path.read_bytes() == before
    assert result["representation"] == "native-docx-ooxml"
    assert result["native_word_pagination_verified"] is False
    assert result["ok"]
    assert result["images"][0]["description"] == "A verified reactor schematic"
    assert result["images"][0]["sha256"] == hashlib.sha256(path.with_suffix(".png").read_bytes()).hexdigest()
    assert result["tables"][0]["rows"] == 85
    assert result["tables"][0]["columns"] == 9
    assert result["math_count"] == result["revision_count"] == 1
    assert result["superscript_count"] == 1
    assert "Thin\u2009space and nonbreaking\u00a0space" in result["extracted_text"]
    assert "A tracked revision." in result["extracted_text"]
    assert len(result["captions"][0]["text"]) > 300


def test_missing_description_after_a_transform_is_visible(tmp_path):
    path = rich_docx(tmp_path / "native.docx", description=None)
    result = inspect_docx(path)
    assert not result["ok"]
    assert any(issue["id"] == "missing-alt-text" for issue in result["issues"])


def test_invalid_docx_is_actionable(tmp_path):
    path = tmp_path / "bad.docx"
    path.write_bytes(b"not a docx")
    with pytest.raises(InputError, match="DOCX|docx"):
        inspect_docx(path)


def test_legacy_image_in_a_footnote_keeps_its_final_alt_and_embedded_identity(tmp_path):
    path = rich_docx(tmp_path / "native.docx")
    with zipfile.ZipFile(path) as package:
        parts = {name: package.read(name) for name in package.namelist()}
    parts["word/footnotes.xml"] = b'''<w:footnotes xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" xmlns:v="urn:schemas-microsoft-com:vml" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"><w:footnote w:id="1"><w:p><w:r><w:pict><v:shape alt="Legacy accessible diagram"><v:imagedata r:id="rId1"/></v:shape></w:pict></w:r></w:p></w:footnote></w:footnotes>'''
    parts["word/_rels/footnotes.xml.rels"] = b'''<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/image" Target="/word/media/image1.png"/></Relationships>'''
    with zipfile.ZipFile(path, "w") as package:
        for name, content in parts.items():
            package.writestr(name, content)
    result = inspect_docx(path)
    image = next(item for item in result["images"] if item["part"] == "word/footnotes.xml")
    assert image["description"] == "Legacy accessible diagram"
    assert image["sha256"] == hashlib.sha256(parts["word/media/image1.png"]).hexdigest()
    assert result["ok"]


@pytest.fixture
def sealed(wongo_project, stub_quarto, monkeypatch):
    import wongo.evidence as evidence
    from wongo.engine import render_project

    monkeypatch.setattr(evidence, "engine_identity", lambda: {"version": "test", "files": {}, "sha256": "stable"})
    render_project(wongo_project, "submission")
    return wongo_project


@pytest.fixture
def fake_preview_tools(monkeypatch):
    import wongo.qa as qa

    content = {}
    monkeypatch.setattr(qa.toolchain, "qa_tool_command", lambda name: [name])
    monkeypatch.setattr(qa.toolchain, "tool_version", lambda *args, **kwargs: "1.2.3")

    def run(command, **kwargs):
        if command[0] == "soffice":
            destination = qa.Path(command[command.index("--outdir") + 1])
            source = qa.Path(command[-1])
            (destination / (source.stem + ".pdf")).write_bytes(b"%PDF-1.4\nsynthetic fixture")
        elif command[0] == "pdfinfo":
            return "Pages: 2\n"
        elif command[0] == "pdftoppm":
            prefix = qa.Path(command[-1])
            for number in (1, 2):
                page = prefix.with_name(prefix.name + f"-{number}.png")
                _png(page)
                page.write_bytes(page.read_bytes() + content.get((prefix.name + ".docx", number), b""))
        return ""

    monkeypatch.setattr(qa, "_run", run)
    return content


def _all_pages(preview):
    return [{"document": name, "page": page["number"]}
            for name, document in preview["result"]["documents"].items() for page in document["pages"]]


def test_default_review_export_does_not_disclose_full_extracted_manuscript_text(sealed, monkeypatch, tmp_path):
    from wongo.engine import render_project
    from wongo.qa import inspect_outputs
    from wongo.review_export import export_review, verify_bundle

    sentinel = "PRIVATE-MANUSCRIPT-BODY-ONLY-716b8a"
    monkeypatch.setenv("WONGO_STUB_QUARTO_STAMP", sentinel)
    render_project(sealed, "submission")
    direct = inspect_docx(sealed / "output" / "main-submission.docx")
    assert sentinel in direct["extracted_text"]
    record = inspect_outputs(sealed)["record"]
    document = record["result"]["documents"]["main-submission.docx"]
    assert "extracted_text" not in document
    assert document["text_identity"] == {
        "sha256": hashlib.sha256(direct["extracted_text"].encode("utf-8")).hexdigest(),
        "characters": len(direct["extracted_text"]),
    }
    destination = tmp_path / "private-review.zip"
    export_review(sealed, destination)
    assert verify_bundle(destination)["ok"]
    with zipfile.ZipFile(destination) as package:
        assert sentinel.encode() not in package.read("review.json")
        assert package.namelist() == ["review.json"]


def test_preview_is_a_separate_artifact_and_does_not_modify_deliverables(sealed, fake_preview_tools):
    from wongo.evidence import safe_path, verify_record
    from wongo.qa import create_preview

    before = {path.name: path.read_bytes() for path in (sealed / "output").iterdir() if path.is_file()}
    result = create_preview(sealed)
    assert result["ok"]
    record = result["record"]
    assert record["result"]["representation"] == "converted-layout-preview"
    assert record["result"]["native_word_verified"] is False
    assert record["result"]["inspection_status"] == "not-recorded"
    assert all(document["page_count"] == 2 for document in record["result"]["documents"].values())
    assert verify_record(sealed, record)["state"] == "fresh"
    assert all(safe_path(sealed, item["path"]).is_file() for item in record["artifacts"])
    assert before == {path.name: path.read_bytes() for path in (sealed / "output").iterdir() if path.is_file()}


def test_preview_storage_cannot_escape_through_a_symlink(sealed, fake_preview_tools, tmp_path):
    from wongo.qa import create_preview

    outside = tmp_path / "external-preview-storage"
    outside.mkdir()
    storage = sealed / "output" / ".wongo-evidence"
    storage.mkdir(exist_ok=True)
    try:
        (storage / "previews").symlink_to(outside, target_is_directory=True)
    except (OSError, NotImplementedError):
        pytest.skip("Symlink creation is unavailable on this host")
    with pytest.raises(InputError, match="escapes"):
        create_preview(sealed)
    assert list(outside.iterdir()) == []


def test_visual_review_needs_human_confirmation_and_reports_untested_pages(sealed, fake_preview_tools):
    from wongo.qa import create_preview, record_visual_review

    preview = create_preview(sealed)["record"]
    selected = _all_pages(preview)[:1]
    with pytest.raises(InputError, match="person"):
        record_visual_review(sealed, "submission", "Ada", selected)
    result = record_visual_review(sealed, "submission", "Ada", selected, confirmed=True)
    assert not result["record"]["result"]["complete"]
    assert len(result["record"]["result"]["untested_pages"]) == 3
    assert result["record"]["result"]["native_word_verified"] is False
    with pytest.raises(InputError, match="absent"):
        record_visual_review(sealed, "submission", "Ada", [{"document": "main-submission.docx", "page": 99}], confirmed=True)


def test_changed_settings_cannot_carry_an_old_inspection(sealed, fake_preview_tools):
    from wongo.qa import carry_visual_review, compare_previews, create_preview, record_visual_review

    first = create_preview(sealed, dpi=96)["record"]
    review = record_visual_review(sealed, "submission", "Ada", _all_pages(first), confirmed=True)["record"]
    second = create_preview(sealed, dpi=144)["record"]
    assert not compare_previews(second, first)["comparable"]
    with pytest.raises(InputError, match="No inspected pages"):
        carry_visual_review(sealed, "submission", second["record_id"], review["record_id"])


def test_page_carry_preserves_attribution_and_exposes_changed_pages(sealed, fake_preview_tools):
    from wongo.qa import carry_visual_review, create_preview, record_visual_review

    first = create_preview(sealed)["record"]
    review = record_visual_review(sealed, "submission", "Ada", _all_pages(first), confirmed=True)["record"]
    fake_preview_tools[("main-submission.docx", 1)] = b"changed scientific content"
    second = create_preview(sealed)["record"]
    carried = carry_visual_review(sealed, "submission", second["record_id"], review["record_id"])["record"]
    result = carried["result"]
    assert result["untested_pages"] == [{"document": "main-submission.docx", "page": 1}]
    assert not result["complete"]
    assert all(page["reviewer"] == "Ada" for page in result["checked_pages"])
    assert all(page["original_record_id"] == review["record_id"] for page in result["checked_pages"])
    assert all(page["original_recorded_at"] == review["recorded_at"] for page in result["checked_pages"])


def test_tampered_preview_artifacts_cannot_be_reviewed(sealed, fake_preview_tools):
    from wongo.qa import create_preview, record_visual_review

    preview = create_preview(sealed)["record"]
    artifact = sealed / preview["artifacts"][0]["path"]
    artifact.write_bytes(b"changed pixels")
    with pytest.raises(InputError, match="stale"):
        record_visual_review(sealed, "submission", "Ada", _all_pages(preview), confirmed=True)


def test_build_change_during_conversion_does_not_save_preview_evidence(sealed, fake_preview_tools, monkeypatch):
    import wongo.qa as qa

    original = qa._run

    def changing(command, **kwargs):
        if command[0] == "soffice":
            source = sealed / "index.qmd"
            source.write_text(source.read_text(encoding="utf-8") + "\nChanged.\n", encoding="utf-8")
        return original(command, **kwargs)

    monkeypatch.setattr(qa, "_run", changing)
    with pytest.raises(InputError, match="stale|changed"):
        qa.create_preview(sealed)
    assert not list((sealed / "output" / ".wongo-evidence").rglob("*.json"))


@pytest.mark.parametrize("dpi", [True, 0, 1000])
def test_preview_resolution_is_validated_before_execution(sealed, dpi):
    from wongo.qa import create_preview
    with pytest.raises(InputError, match="dpi"):
        create_preview(sealed, dpi=dpi)


def test_inspection_after_registered_transform_detects_removed_alt_text(wongo_project, stub_quarto, monkeypatch):
    import wongo.engine as engine
    import wongo.evidence as evidence
    from wongo.qa import inspect_outputs

    monkeypatch.setattr(evidence, "engine_identity", lambda: {"version": "test", "files": {}, "sha256": "stable"})
    original = engine.postprocess_main

    def add_description(path, *args, **kwargs):
        original(path, *args, **kwargs)
        doc = Document(path)
        for element in doc.element.iter():
            if element.tag.endswith("}docPr"):
                element.set("descr", "Original verified description")
        doc.save(path)

    monkeypatch.setattr(engine, "postprocess_main", add_description)
    (wongo_project / "strip.py").write_text(
        "import os\nfrom docx import Document\n"
        "for path in os.environ['WONGO_POSTPROCESS_OUTPUTS'].split(os.pathsep):\n"
        "    doc = Document(path)\n"
        "    for e in doc.element.iter():\n"
        "        if e.tag.endswith('}docPr'): e.attrib.pop('descr', None)\n"
        "    doc.save(path)\n", encoding="utf-8")
    config = wongo_project / "_journal.yml"
    config.write_text(config.read_text(encoding="utf-8") + "postprocess: [strip.py]\n", encoding="utf-8")
    engine.render_project(wongo_project, "submission")
    result = inspect_outputs(wongo_project)
    assert not result["ok"]
    assert any(item["id"] == "missing-alt-text" for item in result["record"]["result"]["documents"]["main-submission.docx"]["issues"])


def test_qa_review_refuses_noninteractive_recording(cli, sealed, monkeypatch):
    from wongo import clitools
    monkeypatch.setattr(clitools, "is_interactive", lambda: False)
    code, stdout, _ = cli("qa", "review", "--project", str(sealed), "--reviewer", "Ada", "--all-pages", "--json")
    import json
    response = json.loads(stdout)
    assert code == 1 and not response["ok"]
    assert response["error"]["kind"] == "input"
    assert "person at a terminal" in response["error"]["message"]


def test_qa_review_confirmation_records_exact_selected_pages(cli, sealed, fake_preview_tools, monkeypatch):
    import json

    from wongo import clitools
    from wongo.qa import create_preview
    create_preview(sealed)
    monkeypatch.setattr(clitools, "is_interactive", lambda: True)
    monkeypatch.setattr("builtins.input", lambda: "1")
    code, stdout, stderr = cli("qa", "review", "--project", str(sealed), "--reviewer", "Ada", "--all-pages", "--json")
    response = json.loads(stdout)
    assert code == 0 and response["ok"]
    assert response["record"]["result"]["complete"]
    assert "native Word pagination stays unverified" in stderr


@pytest.mark.skipif(os.environ.get("WONGO_REAL_QA_TEST") != "1", reason="run in the optional LibreOffice/Poppler QA job")
def test_real_libreoffice_preview_preserves_native_docx_and_records_page_artifacts(
        wongo_project, stub_quarto, monkeypatch):
    import wongo.engine as engine
    import wongo.evidence as evidence
    from wongo.qa import create_preview

    monkeypatch.setattr(evidence, "engine_identity", lambda: {"version": "test", "files": {}, "sha256": "stable"})

    def rich_quarto(project, qmd, out_name, profile, *, stage_dir, **kwargs):
        path = rich_docx(stage_dir / out_name)
        return engine.QuartoRender(path, [])

    monkeypatch.setattr(engine, "quarto_render", rich_quarto)
    result = engine.render_project(wongo_project, "submission")
    before = {path.name: path.read_bytes() for path in result.outputs}
    preview = create_preview(wongo_project, dpi=72)["record"]
    assert all(document["page_count"] >= 2 for document in preview["result"]["documents"].values())
    assert all((wongo_project / item["path"]).is_file() for item in preview["artifacts"])
    assert preview["result"]["native_word_verified"] is False
    assert before == {path.name: path.read_bytes() for path in result.outputs}


def test_native_word_attestation_is_distinct_from_automated_preview_verification(sealed):
    from wongo.qa import record_native_review
    pages = [{"document": "main-submission.docx", "page": 1}]
    documents = {"main-submission.docx": 1, "si-submission.docx": 2}
    with pytest.raises(InputError, match="person"):
        record_native_review(sealed, "submission", "Ada", "Word fixture version", documents, pages)
    result = record_native_review(sealed, "submission", "Ada", "Word fixture version", documents, pages,
                                  confirmed=True)["record"]["result"]
    assert result["representation"] == "native-word-inspection"
    assert result["native_word_attested"] and not result["native_word_verified"]
    assert result["verification_method"] == "local-human-attestation"
    assert not result["identity_authenticated"]
    assert len(result["untested_pages"]) == 2 and not result["complete"]


def test_native_attestation_requires_page_counts_for_all_final_documents(sealed):
    from wongo.qa import record_native_review
    with pytest.raises(InputError, match="every final output"):
        record_native_review(sealed, "submission", "Ada", "test", {"main-submission.docx": 1},
                             [{"document": "main-submission.docx", "page": 1}], confirmed=True)


def test_native_confirmation_refuses_a_different_build(sealed):
    from wongo.qa import record_native_review
    with pytest.raises(InputError, match="changed during confirmation"):
        record_native_review(sealed, "submission", "Ada", "test",
                             {"main-submission.docx": 1, "si-submission.docx": 1},
                             [{"document": "main-submission.docx", "page": 1}],
                             confirmed=True, subject_id="prior-build")


def test_native_cli_refuses_noninteractive_attestation(cli, sealed, monkeypatch):
    import json

    from wongo import clitools
    monkeypatch.setattr(clitools, "is_interactive", lambda: False)
    code, stdout, _ = cli("qa", "native-review", "--project", str(sealed), "--reviewer", "Ada",
                         "--word-version", "test", "--document", "main-submission.docx:1",
                         "--document", "si-submission.docx:1", "--all-pages", "--json")
    response = json.loads(stdout)
    assert code == 1 and response["error"]["kind"] == "input"


def test_renderer_mutating_its_input_copy_is_not_accepted(sealed, fake_preview_tools, monkeypatch):
    import wongo.qa as qa
    original = qa._run

    def mutate_copy(command, **kwargs):
        result = original(command, **kwargs)
        if command[0] == "soffice":
            qa.Path(command[-1]).write_bytes(b"modified by renderer")
        return result

    monkeypatch.setattr(qa, "_run", mutate_copy)
    with pytest.raises(InputError, match="changed its.*input copy"):
        qa.create_preview(sealed)
    assert not list((sealed / "output" / ".wongo-evidence").rglob("*.json"))


def test_native_word_review_cannot_be_carried_as_converted_page_evidence(sealed, fake_preview_tools):
    from wongo.qa import carry_visual_review, create_preview, record_native_review
    native = record_native_review(sealed, "submission", "Ada", "test",
                                  {"main-submission.docx": 1, "si-submission.docx": 1},
                                  [{"document": "main-submission.docx", "page": 1}], confirmed=True)["record"]
    preview = create_preview(sealed)["record"]
    with pytest.raises(InputError, match="Native Word inspection"):
        carry_visual_review(sealed, "submission", preview["record_id"], native["record_id"])


@pytest.mark.parametrize("problem", ["missing-output", "wrong-docx"])
def test_preview_cannot_hide_omitted_or_mismatched_document_evidence(sealed, fake_preview_tools, problem):
    import copy

    from wongo.evidence import make_record, save_record, snapshot
    from wongo.qa import create_preview, record_visual_review

    preview = create_preview(sealed)["record"]
    result = copy.deepcopy(preview["result"])
    if problem == "missing-output":
        del result["documents"]["si-submission.docx"]
    else:
        result["documents"]["main-submission.docx"]["docx_sha256"] = "wrong-source"
    malformed = make_record(snapshot(sealed), "preview", preview["producer"], result,
                            preview["coverage"], artifacts=preview["artifacts"])
    save_record(sealed, malformed)
    with pytest.raises(InputError, match="output document|DOCX identity"):
        record_visual_review(sealed, "submission", "Ada", _all_pages(preview), confirmed=True)
