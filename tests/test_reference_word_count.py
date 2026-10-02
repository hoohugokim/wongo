"""Reference-inclusive gates use this render's main bibliography before promotion."""
from __future__ import annotations

import json

import pytest
import yaml
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from wongo.engine import render_project
from wongo.engine.checks import run_checks
from wongo.errors import GateError


def _reference_project(project, limit=6):
    profile = project / "profiles/demo/profile.yml"
    data = yaml.safe_load(profile.read_text(encoding="utf-8"))
    data["manuscript_types"][0].update(
        word_limit=limit, word_limit_includes_references=True,
        counting_rule="body, abstract and references",
    )
    profile.write_text(yaml.safe_dump(data), encoding="utf-8")
    (project / "index.qmd").write_text("Body @doe2020.\n", encoding="utf-8")
    return project


def _published_files(project):
    return {p.name: p.read_bytes() for p in (project / "output").iterdir() if p.is_file()}


def test_rendered_references_can_fail_submission_without_replacing_previous_outputs(
    stub_quarto, wongo_project, monkeypatch,
):
    project = _reference_project(wongo_project)
    monkeypatch.setenv("WONGO_STUB_QUARTO_REFERENCES", json.dumps(["Doe. Short reference."]))
    render_project(project, "submission")
    previous = _published_files(project)

    # Two source words plus five rendered bibliography words exceed six.
    monkeypatch.setenv("WONGO_STUB_QUARTO_REFERENCES", json.dumps(["Doe. A longer reference here."]))
    with pytest.raises(GateError) as caught:
        render_project(project, "submission")

    check = next(c for c in caught.value.details["checks"] if c.name == "word-limit")
    assert not check.ok and check.level == "HARD"
    assert "7 words" in check.detail
    assert "2 source" in check.detail and "5 reference" in check.detail
    assert _published_files(project) == previous
    assert not list((project / "output").glob(".stage-*"))


def test_exact_limit_uses_main_references_only_and_cli_exposes_the_final_count(
    cli, stub_quarto, wongo_project, monkeypatch,
):
    project = _reference_project(wongo_project)
    monkeypatch.setenv("WONGO_STUB_QUARTO_REFERENCES", json.dumps({
        "index.qmd": ["Doe. Four reference words."],
        "si.qmd": ["Separate supplementary bibliography " * 10],
    }))
    code, out, _ = cli("render", "--project", str(project), "--target", "submission", "--json")
    assert code == 0, out
    result = json.loads(out)
    check = next(c for c in result["checks"] if c["name"] == "word-limit")
    assert check["ok"] and "6 words" in check["detail"]
    assert "2 source + 4 reference" in check["detail"]
    assert "word-limit-references" not in result["warnings"]

    # A source-only check remains honest even with a previous DOCX nearby.
    code, out, _ = cli("check", "--project", str(project), "--strict", "--json")
    assert code == 0
    warn = next(c for c in json.loads(out)["checks"] if c["name"] == "word-limit-references")
    assert not warn["ok"] and "wongo render" in warn["detail"]


def test_collab_can_be_rendered_for_inspection_with_a_failed_reference_gate(
    cli, stub_quarto, wongo_project, monkeypatch,
):
    project = _reference_project(wongo_project)
    monkeypatch.setenv("WONGO_STUB_QUARTO_REFERENCES", json.dumps(["Doe. A longer reference here."]))
    code, out, _ = cli("render", "--project", str(project), "--target", "collab", "--json")
    assert code == 0
    result = json.loads(out)
    assert result["hard_failures"] == ["word-limit"]
    assert (project / "output/main-collab.docx").exists()


def test_bibliography_count_joins_formatted_runs_and_includes_link_text(wongo_project, tmp_path):
    project = _reference_project(wongo_project)
    doc = Document()
    doc.styles.add_style("Bibliography", WD_STYLE_TYPE.PARAGRAPH)
    doc.add_paragraph("Not a bibliography entry.")
    para = doc.add_paragraph(style="Bibliography")
    para.add_run("Doe. ")
    para.add_run("Ti").italic = True
    para.add_run("tle")
    para.add_run().add_tab()
    link = OxmlElement("w:hyperlink")
    link.set(qn("w:anchor"), "example")
    run, text = OxmlElement("w:r"), OxmlElement("w:t")
    text.text = "Linked words."
    run.append(text)
    link.append(run)
    para._p.append(link)
    path = tmp_path / "rendered.docx"
    doc.save(path)

    check = next(c for c in run_checks(project, rendered_main=path) if c.name == "word-limit")
    assert check.ok and "2 source + 4 reference" in check.detail


def test_submission_cannot_treat_an_unrecognized_bibliography_as_zero_words(
    stub_quarto, wongo_project,
):
    project = _reference_project(wongo_project)
    with pytest.raises(GateError) as caught:
        render_project(project, "submission")

    check = next(c for c in caught.value.details["checks"] if c.name == "word-limit-references")
    assert check.level == "HARD" and not check.ok
    assert "Bibliography" in check.detail
    assert not list((project / "output").glob("*.docx"))


def test_source_saved_during_render_cannot_get_a_reference_inclusive_pass(
    stub_quarto, wongo_project, monkeypatch,
):
    project = _reference_project(wongo_project, limit=100)
    monkeypatch.setenv("WONGO_STUB_QUARTO_REFERENCES", json.dumps(["Doe. Short reference."]))
    render_project(project, "submission")
    previous = _published_files(project)
    monkeypatch.setenv("WONGO_STUB_QUARTO_EDIT", "index.qmd")

    with pytest.raises(GateError) as caught:
        render_project(project, "submission")

    check = next(c for c in caught.value.details["checks"] if c.name == "word-limit-references")
    assert not check.ok and "changed during" in check.detail
    assert _published_files(project) == previous


@pytest.mark.parametrize("front,config", [
    ("nocite: '@doe2020'", ""),
    ("", "nocite: '@*'\n"),
    ("", 'abstract: "An abstract cites @doe2020."\n'),
    ("abstract: An abstract cites @doe2020.", ""),
])
def test_metadata_requested_references_must_also_be_visible_in_the_render(
    stub_quarto, wongo_project, front, config,
):
    project = _reference_project(wongo_project, limit=100)
    (project / "index.qmd").write_text(f"---\n{front}\n---\nBody only.\n", encoding="utf-8")
    with (project / "_quarto.yml").open("a", encoding="utf-8") as out:
        out.write(config)
    with pytest.raises(GateError) as caught:
        render_project(project, "submission")
    assert any(c.name == "word-limit-references" and not c.ok
               for c in caught.value.details["checks"])


def test_no_main_citations_can_have_zero_references_even_if_si_has_citations(
    stub_quarto, wongo_project,
):
    project = _reference_project(wongo_project)
    (project / "index.qmd").write_text("Body only.\n", encoding="utf-8")
    (project / "si.qmd").write_text("SI cites @doe2020.\n", encoding="utf-8")
    result = render_project(project, "submission")
    check = next(c for c in result.checks if c.name == "word-limit")
    assert check.ok and "2 source + 0 reference" in check.detail


def test_profile_changed_during_render_cannot_disable_the_reference_gate(
    stub_quarto, wongo_project, monkeypatch,
):
    project = _reference_project(wongo_project)
    monkeypatch.setenv("WONGO_STUB_QUARTO_REFERENCES", json.dumps(["Doe. Short reference."]))
    monkeypatch.setenv("WONGO_STUB_QUARTO_EDIT", "profiles/demo/profile.yml")
    monkeypatch.setenv("WONGO_STUB_QUARTO_EDIT_TEXT",
                      '\nmanuscript_types: [{type: article, word_limit: 1000}]\n')
    with pytest.raises(GateError) as caught:
        render_project(project, "submission")
    check = next(c for c in caught.value.details["checks"] if c.name == "word-limit-references")
    assert not check.ok and "changed during" in check.detail
    assert not list((project / "output").glob("*.docx"))
