"""Explicit SI references use final bookmarked captions, never semantic label digits."""
from __future__ import annotations

import hashlib
import io
import json
import struct
import zipfile
import zlib

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from wongo.engine import QuartoRender, output_freshness, read_manifest, render_project
from wongo.engine.checks import run_checks
from wongo.engine.si_refs import enabled, resolve, source_requests
from wongo.errors import ConfigError, GateError, InputError


def _link(paragraph, label="tbl-beta", text="SI table"):
    rel = paragraph.part.relate_to(f"wongo-si:{label}",
                                  "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
                                  is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), rel)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    props.append(OxmlElement("w:b"))
    run.append(props)
    node = OxmlElement("w:t")
    node.text = text
    run.append(node)
    link.append(run)
    paragraph._p.append(link)
    return link


def _main(path, labels=("tbl-beta",), *, rich=False):
    doc = Document()
    paragraph = doc.add_paragraph("See ")
    for label in labels:
        link = _link(paragraph, label)
        if rich:
            link[0].append(OxmlElement("w:fldChar"))
        paragraph.add_run(" in the Supporting Information. ")
    equation = doc.add_paragraph("Equation remains: ")
    equation._p.append(OxmlElement("m:oMath"))
    doc.add_table(rows=1, cols=1).cell(0, 0).text = "Untouched scientific value 42"
    def chunk(kind, data):
        return struct.pack(">I", len(data)) + kind + data + struct.pack(">I", zlib.crc32(kind + data) & 0xFFFFFFFF)
    png = (b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", struct.pack(">IIBBBBB", 2, 1, 8, 2, 0, 0, 0))
           + chunk(b"IDAT", zlib.compress(b"\x00\x00\x80\x40\x00\x80\x40")) + chunk(b"IEND", b""))
    doc.add_picture(io.BytesIO(png))
    doc.save(path)


def _si(path, items=(("tbl-beta", "Table S7: Observations."),), *, styled=True):
    doc = Document()
    for index, (label, caption) in enumerate(items):
        cell = doc.add_table(rows=1, cols=1).cell(0, 0)
        start, end = OxmlElement("w:bookmarkStart"), OxmlElement("w:bookmarkEnd")
        start.set(qn("w:name"), label)
        start.set(qn("w:id"), str(index))
        end.set(qn("w:id"), str(index))
        cell._tc.append(start)
        cell.add_paragraph(caption, style="Caption" if styled else "Normal")
        data = cell.add_table(rows=2, cols=2)
        data.cell(0, 0).text = "Table S999: Scientific data, not a caption."
        cell._tc.append(end)
    doc.save(path)


def _parts(path):
    with zipfile.ZipFile(path) as package:
        return {name: package.read(name) for name in package.namelist()}


@pytest.mark.parametrize("styled", [True, False])
def test_exact_final_caption_and_other_xml_parts_are_preserved(tmp_path, styled):
    main, si = tmp_path / "main.docx", tmp_path / "si.docx"
    _main(main, ("tbl-beta", "fig-alpha", "tbl-beta"))
    _si(si, (("fig-alpha", "Figure S 3. Diagram."), ("tbl-beta", "Table S7: Observations.")), styled=styled)
    before, si_before = _parts(main), si.read_bytes()
    record = resolve(main, si, active=True)
    assert [item["display"] for item in record["references"]] == ["Table S7", "Figure S3", "Table S7"]
    after = _parts(main)
    assert set(after) == set(before)
    assert {name for name in before if before[name] != after[name]} == {
        "word/document.xml", "word/_rels/document.xml.rels"}
    assert b"Untouched scientific value 42" in after["word/document.xml"]
    assert b"m:oMath" in after["word/document.xml"]
    assert any(name.startswith("word/media/") for name in after)
    assert b"wongo-si:" not in after["word/_rels/document.xml.rels"]
    assert si.read_bytes() == si_before
    assert record["si_sha256"] == hashlib.sha256(si_before).hexdigest()
    assert record["main_after"] == hashlib.sha256(main.read_bytes()).hexdigest()
    assert resolve(main, si, active=True)["references"] == []  # no second mutation
    assert _parts(main) == after


@pytest.mark.parametrize("items", [(), (("tbl-beta", "Table S1: One."), ("tbl-beta", "Table S2: Two.")),
                                    (("tbl-beta", "Figure S1: Wrong type."),),
                                    (("tbl-beta", "Unnumbered caption."),)])
def test_missing_ambiguous_or_unsupported_final_destinations_refuse(tmp_path, items):
    main, si = tmp_path / "main.docx", tmp_path / "si.docx"
    _main(main)
    _si(si, items)
    before = main.read_bytes()
    with pytest.raises(InputError):
        resolve(main, si, active=True)
    assert main.read_bytes() == before


@pytest.mark.parametrize("label", ["eq-science", "tbl-beta?version=old", "tbl-beta#caption", "../tbl-beta"])
def test_only_supported_exact_float_labels_are_accepted(tmp_path, label):
    main = tmp_path / "main.docx"
    _main(main, (label,))
    with pytest.raises(InputError, match="Invalid SI reference"):
        resolve(main, None, active=True)


def test_undeclared_or_rich_requests_refuse_without_mutation(tmp_path):
    main = tmp_path / "main.docx"
    _main(main)
    before = main.read_bytes()
    with pytest.raises(InputError, match="si_references: true"):
        resolve(main, None, active=False)
    assert main.read_bytes() == before
    _main(main, rich=True)
    with pytest.raises(InputError, match="cannot be replaced safely"):
        resolve(main, None, active=True)


def test_literal_code_comments_and_metadata_are_not_source_requests():
    text = ('---\ntitle: "[SI](wongo-si:tbl-meta)"\n---\n'
            '<!-- [SI](wongo-si:tbl-comment) -->\n'
            '`[SI](wongo-si:tbl-code)`\n\n'
            '```markdown\n[SI](wongo-si:tbl-example)\n```\n\n'
            '[SI](wongo-si:tbl-real)\n')
    assert source_requests(text) == [("tbl-real", 11)]
    assert source_requests(r"\[literal](wongo-si:tbl-example)") == []


@pytest.mark.parametrize("value", ["true", 1, [], None])
def test_opt_in_is_a_boolean(value):
    with pytest.raises(ConfigError, match="true or false"):
        enabled({"si_references": value})


def _configure(project):
    path = project / "_journal.yml"
    path.write_text(path.read_text(encoding="utf-8") + "si_references: true\n", encoding="utf-8")
    source = project / "index.qmd"
    source.write_text(source.read_text(encoding="utf-8") + "\nSee [SI](wongo-si:tbl-beta).\n", encoding="utf-8")


@pytest.fixture
def fake_render(monkeypatch):
    import wongo.engine as engine

    def render(project, qmd, out_name, profile, *, stage_dir, **kwargs):
        path = stage_dir / out_name
        if qmd == "si.qmd":
            _si(path)
        else:
            _main(path)
        return QuartoRender(path, [])
    monkeypatch.setattr(engine, "quarto_render", render)


@pytest.mark.parametrize("target", ["collab", "submission"])
def test_resolution_after_declared_transforms_is_sealed_in_manifest(
        wongo_project, stub_quarto, fake_render, target):
    _configure(wongo_project)
    script = wongo_project / "finish.py"
    script.write_text(
        'import os,zipfile\nfrom pathlib import Path\n'
        'p=next(Path(raw) for raw in os.environ["WONGO_POSTPROCESS_OUTPUTS"].split(os.pathsep) if Path(raw).name.startswith("si-"))\n'
        'with zipfile.ZipFile(p) as z: parts={n:z.read(n) for n in z.namelist()}\n'
        'parts["word/document.xml"]=parts["word/document.xml"].replace(b"Table S7",b"Table S8")\n'
        'with zipfile.ZipFile(p,"w") as z:\n'
        ' for name,raw in parts.items(): z.writestr(name,raw)\n', encoding="utf-8")
    cfg = wongo_project / "_journal.yml"
    cfg.write_text(cfg.read_text(encoding="utf-8") + "postprocess: [finish.py]\n", encoding="utf-8")
    source_before = (wongo_project / "index.qmd").read_bytes()
    result = render_project(wongo_project, target)
    assert not result.hard_failures
    assert "Table S8" in Document(wongo_project / "output" / f"main-{target}.docx").paragraphs[0].text
    entry = read_manifest(wongo_project)["targets"][target]
    assert entry["si_references"]["references"][0]["display"] == "Table S8"
    assert entry["si_references"]["main_after"] == entry["outputs"][f"main-{target}.docx"]
    assert entry["si_references"]["si_sha256"] == entry["outputs"][f"si-{target}.docx"]
    assert output_freshness(wongo_project, target)["state"] == "fresh"
    assert (wongo_project / "index.qmd").read_bytes() == source_before


@pytest.mark.parametrize("target", ["collab", "submission"])
def test_missing_final_label_preserves_both_old_outputs_and_manifest(
        wongo_project, stub_quarto, fake_render, monkeypatch, target):
    _configure(wongo_project)
    render_project(wongo_project, target)
    before = {p.name: p.read_bytes() for p in (wongo_project / "output").iterdir() if p.is_file()}
    import wongo.engine.si_refs as refs

    original = refs._caption_labels
    monkeypatch.setattr(refs, "_caption_labels", lambda raw, requested: original(raw, {"tbl-missing"}))
    with pytest.raises(InputError, match="tbl-missing"):
        render_project(wongo_project, target)
    after = {p.name: p.read_bytes() for p in (wongo_project / "output").iterdir() if p.is_file()}
    assert before == after
    assert not list((wongo_project / "output").glob(".stage-*"))


def test_native_cross_document_reference_is_still_invalid(wongo_project):
    _configure(wongo_project)
    si = wongo_project / "si.qmd"
    si.write_text(si.read_text(encoding="utf-8") + "\nTable: Values {#tbl-beta}\n", encoding="utf-8")
    main = wongo_project / "index.qmd"
    main.write_text(main.read_text(encoding="utf-8") + "\nWrong native @tbl-beta.\n", encoding="utf-8")
    check = next(check for check in run_checks(wongo_project) if check.name == "crossrefs")
    assert not check.ok and "index.qmd: @tbl-beta" in check.detail


def test_missing_opt_in_is_a_source_failure_for_both_targets(wongo_project, stub_quarto):
    source = wongo_project / "index.qmd"
    source.write_text(source.read_text(encoding="utf-8") + "\n[SI](wongo-si:tbl-beta)\n", encoding="utf-8")
    check = next(check for check in run_checks(wongo_project) if check.name == "si-references")
    assert not check.ok
    for target in ("collab", "submission"):
        with pytest.raises(GateError, match="explicit SI reference"):
            render_project(wongo_project, target)
    assert not stub_quarto.renders()


def test_explicit_main_callout_satisfies_si_citation_without_native_pooling(wongo_project):
    _configure(wongo_project)
    si = wongo_project / "si.qmd"
    si.write_text(si.read_text(encoding="utf-8") + "\nTable: Observations {#tbl-beta}\n", encoding="utf-8")
    checks = run_checks(wongo_project)
    assert next(check for check in checks if check.name == "crossrefs").ok
    assert next(check for check in checks if check.name == "uncited").ok


def test_header_requests_cannot_leak_as_broken_links(tmp_path):
    main = tmp_path / "main.docx"
    doc = Document()
    _link(doc.sections[0].header.paragraphs[0])
    doc.save(main)
    with pytest.raises(InputError, match="header1.xml are unsupported"):
        resolve(main, None, active=True)


def test_old_render_without_si_resolution_identity_is_stale(wongo_project, stub_quarto, fake_render):
    _configure(wongo_project)
    render_project(wongo_project, "submission")
    manifest = read_manifest(wongo_project)
    del manifest["targets"]["submission"]["si_references"]
    path = wongo_project / "output/.wongo-manifest.json"
    path.write_text(json.dumps(manifest), encoding="utf-8")
    result = output_freshness(wongo_project, "submission")
    assert result["state"] == "stale" and "si reference identity" in result["changed"]


def test_explicit_refs_do_not_hide_collab_native_warning(wongo_project, stub_quarto, fake_render, monkeypatch):
    import wongo.engine as engine

    _configure(wongo_project)
    original = engine.quarto_render
    def warning(*args, **kwargs):
        result = original(*args, **kwargs)
        return QuartoRender(result.path, ["Unable to resolve crossref @tbl-native-dynamic"])
    monkeypatch.setattr(engine, "quarto_render", warning)
    render_project(wongo_project, "collab")
    entry = read_manifest(wongo_project)["targets"]["collab"]
    assert "crossrefs-rendered" in entry["checks_failed"]
    assert entry["si_references"]["references"][0]["display"] == "Table S7"


@pytest.mark.parametrize("target", ["collab", "submission"])
def test_source_drift_during_explicit_ref_render_still_refuses(wongo_project, stub_quarto, fake_render, monkeypatch, target):
    import wongo.engine as engine

    _configure(wongo_project)
    render_project(wongo_project, target)
    before = {p.name: p.read_bytes() for p in (wongo_project / "output").iterdir() if p.is_file()}
    original = engine.quarto_render
    def edit(*args, **kwargs):
        result = original(*args, **kwargs)
        source = wongo_project / "index.qmd"
        source.write_text(source.read_text(encoding="utf-8") + "\nConcurrent author save.\n", encoding="utf-8")
        return result
    monkeypatch.setattr(engine, "quarto_render", edit)
    with pytest.raises(InputError, match="changed during rendering"):
        render_project(wongo_project, target)
    assert {p.name: p.read_bytes() for p in (wongo_project / "output").iterdir() if p.is_file()} == before


def test_deleted_reference_does_not_invent_a_visible_requirement(tmp_path):
    main = tmp_path / "main.docx"
    doc = Document()
    p = doc.add_paragraph()
    link = _link(p, "tbl-removed")
    p._p.remove(link)
    deletion = OxmlElement("w:del")
    deletion.append(link)
    p._p.append(deletion)
    doc.save(main)
    before = main.read_bytes()
    assert resolve(main, None, active=True)["references"] == []
    assert main.read_bytes() == before


def test_caption_style_inside_a_data_cell_is_not_the_float_caption(tmp_path):
    main, si = tmp_path / "main.docx", tmp_path / "si.docx"
    _main(main)
    _si(si, items=(("tbl-beta", "Unnumbered caption."),))
    doc = Document(si)
    doc.tables[0].cell(0, 0).tables[0].cell(0, 0).paragraphs[0].style = "Caption"
    doc.save(si)
    with pytest.raises(InputError, match="one numbered Figure/Table caption"):
        resolve(main, si, active=True)
