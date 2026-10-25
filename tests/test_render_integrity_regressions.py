"""Regressions for the final render boundary and recorded input identity."""
from __future__ import annotations

import json
import subprocess
import sys
from pathlib import Path

import pytest
import yaml
from docx import Document

from wongo.engine import output_freshness, render_project, source_fingerprint
from wongo.errors import ConfigError, GateError, InputError
from wongo.styles import styles_dir


def _profile(project: Path, **values) -> None:
    path = project / "profiles" / "demo" / "profile.yml"
    profile = yaml.safe_load(path.read_text(encoding="utf-8"))
    profile.update(values)
    path.write_text(yaml.safe_dump(profile), encoding="utf-8")


def _script(project: Path, body: str) -> Path:
    path = project / "finish.py"
    path.write_text(body, encoding="utf-8")
    config = project / "_journal.yml"
    config.write_text(config.read_text(encoding="utf-8") + "postprocess: [finish.py]\n",
                      encoding="utf-8")
    return path


def _delivered(project: Path) -> dict[str, bytes]:
    return {path.name: path.read_bytes() for path in (project / "output").iterdir()
            if path.name in {"main-submission.docx", "si-submission.docx", ".wongo-manifest.json"}}


def test_final_reference_gate_validates_transformed_bytes_and_keeps_prior_delivery(
        stub_quarto, wongo_project, monkeypatch):
    project = wongo_project
    (project / "index.qmd").write_text("one two\n", encoding="utf-8")
    _profile(project, manuscript_types=[{
        "type": "article", "word_limit": 6, "counting_rule": "references included",
        "word_limit_includes_references": True,
    }])
    monkeypatch.setenv("WONGO_STUB_QUARTO_REFERENCES", '["three four five"]')
    render_project(project, "submission")
    previous = _delivered(project)
    _script(project, "import os\nfrom docx import Document\n"
            "path = os.environ['WONGO_POSTPROCESS_OUTPUTS'].split(os.pathsep)[0]\n"
            "doc = Document(path)\n"
            "doc.add_paragraph('six seven eight nine ten', style='Bibliography')\n"
            "doc.save(path)\n")

    with pytest.raises(GateError) as caught:
        render_project(project, "submission")

    assert any(not check.ok and "10 words" in check.detail
               for check in caught.value.details["checks"] if check.name == "word-limit")
    assert _delivered(project) == previous
    assert not (project / "output" / ".stage-submission").exists()


def test_postprocess_stdout_keeps_json_parseable(stub_quarto, wongo_project):
    _script(wongo_project, "print('HOOK_NOISE')\n")
    completed = subprocess.run(
        [sys.executable, "-m", "wongo.cli", "render", "--target", "collab",
         "--project", str(wongo_project), "--json"],
        capture_output=True, text=True, encoding="utf-8", check=False,
    )
    assert completed.returncode == 0, completed.stderr
    assert json.loads(completed.stdout)["ok"] is True
    assert "HOOK_NOISE" in completed.stderr


def test_uppercase_python_postprocess_uses_the_python_interpreter(stub_quarto, wongo_project):
    script = _script(wongo_project, "# Python script without execute permissions\n")
    upper = script.with_name("FINISH.PY")
    script.rename(upper)
    config = wongo_project / "_journal.yml"
    config.write_text(config.read_text(encoding="utf-8").replace("finish.py", "FINISH.PY"),
                      encoding="utf-8")
    render_project(wongo_project, "collab")
    assert output_freshness(wongo_project, "collab")["state"] == "fresh"


@pytest.mark.parametrize("target", ["collab", "submission"])
@pytest.mark.parametrize("asset", ["csl", "reference_doc", "style"])
def test_a_mid_render_asset_change_refuses_promotion(
        stub_quarto, wongo_project, tmp_path, monkeypatch, target, asset):
    project = wongo_project
    if asset == "style":
        styles = tmp_path / "styles"
        styles.mkdir()
        path = styles / "default.yml"
        path.write_bytes((styles_dir() / "default.yml").read_bytes())
        monkeypatch.setenv("WONGO_STYLES_DIR", str(styles))
    else:
        path = project / "profiles" / "demo" / ("example.csl" if asset == "csl" else "reference.docx")
        if asset == "reference_doc":
            Document().save(path)
        else:
            path.write_text("<style/>\n", encoding="utf-8")
        _profile(project, **{asset: path.name})
    monkeypatch.setenv("WONGO_STUB_QUARTO_EDIT", str(path))
    monkeypatch.setenv("WONGO_STUB_QUARTO_EDIT_TEXT", "\n# changed\n")

    with pytest.raises(InputError, match="changed during rendering"):
        render_project(project, target)

    assert not (project / "output" / f"main-{target}.docx").exists()
    assert not (project / "output" / ".wongo-manifest.json").exists()


def test_an_explicit_style_override_stays_fresh(stub_quarto, wongo_project):
    render_project(wongo_project, "collab", style_override="kist-wcr")
    assert output_freshness(wongo_project, "collab")["state"] == "fresh"
    config = wongo_project / "_journal.yml"
    config.write_text(config.read_text(encoding="utf-8").replace("style: default", "style: kist-wcr"),
                      encoding="utf-8")
    assert output_freshness(wongo_project, "collab")["state"] == "stale"


def test_an_environment_style_change_makes_a_non_overridden_render_stale(
        stub_quarto, wongo_project, monkeypatch):
    config = wongo_project / "_journal.yml"
    config.write_text(config.read_text(encoding="utf-8").replace("style: default\n", ""),
                      encoding="utf-8")
    monkeypatch.setenv("WONGO_STYLE", "default")
    render_project(wongo_project, "collab")
    monkeypatch.setenv("WONGO_STYLE", "kist-wcr")
    assert output_freshness(wongo_project, "collab")["state"] == "stale"


def test_highlights_add_edit_and_remove_change_the_fingerprint(stub_quarto, wongo_project):
    path = wongo_project / "highlights.md"
    absent = source_fingerprint(wongo_project)
    path.write_text("- one\n- two\n- three\n", encoding="utf-8")
    present = source_fingerprint(wongo_project)
    assert present != absent
    render_project(wongo_project, "collab")
    path.write_text("- changed\n- two\n- three\n", encoding="utf-8")
    assert output_freshness(wongo_project, "collab")["state"] == "stale"
    path.unlink()
    assert source_fingerprint(wongo_project) == absent


@pytest.mark.parametrize("target", ["collab", "submission"])
def test_a_mid_render_highlights_change_refuses_promotion(
        stub_quarto, wongo_project, monkeypatch, target):
    path = wongo_project / "highlights.md"
    path.write_text("- one\n- two\n- three\n", encoding="utf-8")
    monkeypatch.setenv("WONGO_STUB_QUARTO_EDIT", str(path))
    with pytest.raises(InputError, match="changed during rendering"):
        render_project(wongo_project, target)


@pytest.mark.parametrize("target", ["collab", "submission"])
def test_a_postprocess_profile_change_refuses_promotion(stub_quarto, wongo_project, target):
    _script(wongo_project, "import os\nfrom pathlib import Path\n"
            "path = Path(os.environ['WONGO_PROJECT']) / 'profiles/demo/profile.yml'\n"
            "path.write_text(path.read_text() + '\\n# changed\\n')\n")
    with pytest.raises(InputError, match="changed during rendering"):
        render_project(wongo_project, target)


def test_final_checks_keep_quarto_crossref_failures(stub_quarto, wongo_project, monkeypatch):
    _profile(wongo_project, manuscript_types=[{
        "type": "article", "word_limit": 500, "counting_rule": "references included",
        "word_limit_includes_references": True,
    }])
    monkeypatch.setenv("WONGO_STUB_QUARTO_REFERENCES", '["one two three"]')
    monkeypatch.setenv("WONGO_STUB_QUARTO_STDOUT", "WARNING: Unable to resolve crossref @fig-missing")
    _script(wongo_project, "# registered identity transform\n")
    result = render_project(wongo_project, "collab")
    assert any(check.name == "crossrefs-rendered" and not check.ok for check in result.checks)
    entry = json.loads(result.manifest.read_text(encoding="utf-8"))["targets"]["collab"]
    assert entry["checks_failed"] == ["crossrefs-rendered"]
    assert output_freshness(wongo_project, "collab")["state"] == "fresh"


def test_engine_reports_diagnostics_without_printing(stub_quarto, wongo_project, capfd):
    _script(wongo_project, "import sys\nprint('hook out')\nprint('hook err', file=sys.stderr)\n")
    events = []
    render_project(wongo_project, "collab", on_event=lambda kind, payload: events.append((kind, payload)))
    captured = capfd.readouterr()
    assert captured.out == captured.err == ""
    diagnostics = "\n".join(payload["text"] for kind, payload in events if kind == "diagnostic")
    assert "hook out" in diagnostics and "hook err" in diagnostics
    assert "Output created:" in diagnostics


@pytest.mark.parametrize("command", [
    'python "scripts/keep captions.py"',
    '/usr/bin/python3 "scripts/keep captions.py"',
    '"scripts/keep captions.py"',
])
def test_quarto_hooks_track_quoted_script_paths_not_interpreters(wongo_project, command):
    script = wongo_project / "scripts" / "keep captions.py"
    script.parent.mkdir()
    script.write_text("# helper\n", encoding="utf-8")
    config = wongo_project / "_quarto.yml"
    config.write_text(yaml.safe_dump({"project": {"pre-render": [command]}}), encoding="utf-8")
    before = source_fingerprint(wongo_project)
    assert before["scripts/keep captions.py"] != "missing"
    assert "/usr/bin/python3" not in before
    script.write_text("# changed\n", encoding="utf-8")
    assert before != source_fingerprint(wongo_project)


def test_quarto_hook_tokenization_preserves_windows_backslashes(wongo_project):
    from wongo.engine.inputs import _hook_tokens

    config = wongo_project / "_quarto.yml"
    config.write_text(yaml.safe_dump({"project": {
        "post-render": [r'C:\Python313\python.exe "scripts\keep captions.py"'],
    }}), encoding="utf-8")
    assert _hook_tokens(wongo_project) == [r"scripts\keep captions.py"]


def test_an_asset_changed_by_initial_check_callback_refuses_promotion(
        stub_quarto, wongo_project, tmp_path, monkeypatch):
    styles = tmp_path / "styles"
    styles.mkdir()
    path = styles / "default.yml"
    path.write_bytes((styles_dir() / "default.yml").read_bytes())
    monkeypatch.setenv("WONGO_STYLES_DIR", str(styles))

    def on_event(kind, payload):
        if kind == "checks":
            path.write_text(path.read_text(encoding="utf-8") + "\n# changed\n", encoding="utf-8")

    with pytest.raises(InputError, match="changed during rendering"):
        render_project(wongo_project, "collab", on_event=on_event)
    assert stub_quarto.renders() == []


@pytest.mark.parametrize("action", ["path.write_bytes(b'not a DOCX')", "path.unlink()"])
def test_postprocess_invalid_or_missing_docx_keeps_previous_delivery(
        stub_quarto, wongo_project, action):
    render_project(wongo_project, "submission")
    previous = _delivered(wongo_project)
    _script(wongo_project, "import os\nfrom pathlib import Path\n"
            "for raw in os.environ['WONGO_POSTPROCESS_OUTPUTS'].split(os.pathsep):\n"
            "    path = Path(raw)\n"
            f"    {action}\n")
    with pytest.raises(InputError, match="finish.py"):
        render_project(wongo_project, "submission")
    assert _delivered(wongo_project) == previous


@pytest.mark.parametrize("declaration", ["file", "directory"])
def test_explicit_external_dependencies_authorize_the_same_quarto_hook(
        wongo_project, tmp_path, declaration):
    outside = tmp_path / "helpers"
    outside.mkdir()
    script = outside / "keep.py"
    script.write_text("# helper\n", encoding="utf-8")
    quarto = wongo_project / "_quarto.yml"
    quarto.write_text(yaml.safe_dump({"project": {"post-render": "python ../helpers/keep.py"}}),
                      encoding="utf-8")
    journal = wongo_project / "_journal.yml"
    cfg = yaml.safe_load(journal.read_text(encoding="utf-8"))
    cfg["dependencies"] = [{
        "path": "../helpers/keep.py" if declaration == "file" else "../helpers",
        "external": True,
    }]
    journal.write_text(yaml.safe_dump(cfg), encoding="utf-8")
    assert script.resolve().as_posix() in source_fingerprint(wongo_project)


def test_an_external_quarto_hook_needs_an_explicit_declaration(wongo_project, tmp_path):
    script = tmp_path / "outside.py"
    script.write_text("# helper\n", encoding="utf-8")
    (wongo_project / "_quarto.yml").write_text(yaml.safe_dump({
        "project": {"post-render": "python ../outside.py"},
    }), encoding="utf-8")
    with pytest.raises(ConfigError, match="external: true"):
        source_fingerprint(wongo_project)


@pytest.mark.parametrize("name", ["pre-hook", "pre-hook.custom"])
def test_interpreter_hooks_track_extensionless_and_custom_suffix_scripts(wongo_project, name):
    script = wongo_project / name
    script.write_text("# helper\n", encoding="utf-8")
    (wongo_project / "_quarto.yml").write_text(yaml.safe_dump({
        "project": {"pre-render": f"/usr/bin/python3 -W ignore -X utf8 {name}"},
    }), encoding="utf-8")
    before = source_fingerprint(wongo_project)
    assert name in before
    assert "ignore" not in before and "utf8" not in before
    script.write_text("# changed\n", encoding="utf-8")
    assert before != source_fingerprint(wongo_project)


def test_missing_interpreter_script_keeps_a_named_missing_fingerprint(wongo_project):
    (wongo_project / "_quarto.yml").write_text(yaml.safe_dump({
        "project": {"post-render": "python -- missing-helper"},
    }), encoding="utf-8")
    assert source_fingerprint(wongo_project)["missing-helper"] == "missing"


def test_direct_existing_hook_tracks_a_custom_suffix(wongo_project):
    (wongo_project / "pre-hook.custom").write_text("# helper\n", encoding="utf-8")
    (wongo_project / "_quarto.yml").write_text(yaml.safe_dump({
        "project": {"pre-render": "pre-hook.custom"},
    }), encoding="utf-8")
    assert "pre-hook.custom" in source_fingerprint(wongo_project)


@pytest.mark.parametrize("command", [
    'python -c "print(123)"',
    'python -c"print(123)"',
    'python -m package.module',
    'node --eval "console.log(123)"',
])
def test_hook_inline_code_and_modules_are_not_script_paths(wongo_project, command):
    from wongo.engine.inputs import _hook_tokens

    (wongo_project / "_quarto.yml").write_text(yaml.safe_dump({
        "project": {"pre-render": command},
    }), encoding="utf-8")
    assert _hook_tokens(wongo_project) == []
