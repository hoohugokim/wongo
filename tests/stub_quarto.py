"""Stand-in for the quarto CLI in tests: no Quarto, pandoc or R needed.

- ``render <qmd> --to docx --output NAME ...`` writes a small DOCX where Quarto
  would: the project output-dir from _quarto.yml, else next to the input.
- ``pandoc --track-changes=all <docx> -t markdown`` prints the UTF-8 file named
  by WONGO_STUB_PANDOC_MD_FILE.
- ``--version`` prints 1.10.18.

Every invocation is appended to WONGO_STUB_QUARTO_LOG as a JSON line, and
WONGO_STUB_QUARTO_FAIL=<qmd> makes rendering that file exit 1.
"""
from __future__ import annotations

import json
import os
import struct
import sys
import zlib
from pathlib import Path


def _log(argv: list[str]) -> None:
    log = os.environ.get("WONGO_STUB_QUARTO_LOG")
    if log:
        with open(log, "a", encoding="utf-8") as fh:
            fh.write(json.dumps({"argv": argv, "cwd": os.getcwd()}, ensure_ascii=False) + "\n")


def _output_dir(project: Path) -> Path:
    config = project / "_quarto.yml"
    if config.exists():
        import yaml

        data = yaml.safe_load(config.read_text(encoding="utf-8-sig")) or {}
        out = (data.get("project") or {}).get("output-dir")
        if out:
            return project / out
    return project


def _png(path: Path, width: int = 8, height: int = 6) -> Path:
    def chunk(tag: bytes, data: bytes) -> bytes:
        body = tag + data
        return struct.pack(">I", len(data)) + body + struct.pack(">I", zlib.crc32(body) & 0xFFFFFFFF)

    raw = b"".join(b"\x00" + b"\x40\x80\xc0" * width for _ in range(height))
    path.write_bytes(b"\x89PNG\r\n\x1a\n"
                     + chunk(b"IHDR", struct.pack(">IIBBBBB", width, height, 8, 2, 0, 0, 0))
                     + chunk(b"IDAT", zlib.compress(raw))
                     + chunk(b"IEND", b""))
    return path


def _docx(path: Path, qmd: str) -> None:
    from docx import Document

    si = qmd.startswith("si")
    prefix = "S" if si else ""
    path.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    doc.add_paragraph("Stub manuscript", style="Title")
    doc.add_paragraph(f"Body rendered from {qmd}.")
    figure = doc.add_table(rows=1, cols=1)  # Quarto wraps a labelled figure in a 1x1 table
    cell = figure.cell(0, 0)
    cell.paragraphs[0].add_run().add_picture(str(_png(path.parent / f".{path.stem}.png")))
    cell.add_paragraph(f"Figure {prefix}1: A stub figure.")
    table = doc.add_table(rows=1, cols=1)  # ...and a labelled table float too
    table.cell(0, 0).paragraphs[0].text = f"Table {prefix}1: A stub table."
    table.cell(0, 0).add_table(rows=2, cols=2)
    doc.add_paragraph("Closing sentence.")
    doc.save(str(path))
    (path.parent / f".{path.stem}.png").unlink()


def main(argv: list[str]) -> int:
    _log(argv)
    if argv[:1] == ["--version"]:
        print("1.10.18")
        return 0
    if argv[:1] == ["render"]:
        qmd = argv[1]
        name = argv[argv.index("--output") + 1]
        if os.environ.get("WONGO_STUB_QUARTO_FAIL") == qmd:
            print(f"stub quarto: failing {qmd} on request", file=sys.stderr)
            return 1
        _docx(_output_dir(Path.cwd()) / name, qmd)
        print(f"Output created: {name}", file=sys.stderr)
        return 0
    if argv[:1] == ["pandoc"]:
        source = os.environ.get("WONGO_STUB_PANDOC_MD_FILE")
        if source:
            sys.stdout.buffer.write(Path(source).read_bytes())
        return 0
    print(f"stub quarto: unsupported arguments {argv}", file=sys.stderr)
    return 2


if __name__ == "__main__":
    raise SystemExit(main(sys.argv[1:]))
