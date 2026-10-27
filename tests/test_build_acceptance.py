"""Acceptance checks across the render, CLI, and declared-transform boundaries."""
from __future__ import annotations

import json
import os
import subprocess
import sys
import zipfile

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree

from wongo.engine import output_freshness, read_manifest, render_project


@pytest.mark.parametrize("failure", [False, True])
def test_postprocess_diagnostics_keep_cli_json_parseable(
        stub_quarto, wongo_project, failure):
    script = wongo_project / "finish.py"
    script.write_text(
        "import sys\n"
        "print('hook stdout')\n"
        "print('hook stderr', file=sys.stderr)\n"
        f"sys.exit({int(failure)})\n",
        encoding="utf-8",
    )
    config = wongo_project / "_journal.yml"
    config.write_text(config.read_text(encoding="utf-8")
                      + "postprocess: [finish.py]\n", encoding="utf-8")

    result = subprocess.run(
        [sys.executable, "-m", "wongo.cli", "render", "--target", "collab",
         "--project", str(wongo_project), "--json"],
        capture_output=True, text=True, encoding="utf-8", check=False,
    )

    response = json.loads(result.stdout)
    assert result.returncode == int(failure)
    assert response["ok"] is not failure
    assert "hook stdout" in result.stderr
    assert "hook stderr" in result.stderr


def _parts(path):
    with zipfile.ZipFile(path) as package:
        return {name: package.read(name) for name in package.namelist()}


def test_registered_idempotent_transform_preserves_unrelated_docx_content(
        stub_quarto, wongo_project, monkeypatch, tmp_path):
    import wongo.engine as engine

    original = engine.postprocess_main
    before = {}

    def seed_rich_content(path, *args, **kwargs):
        original(path, *args, **kwargs)
        doc = Document(path)
        inserted = OxmlElement("w:ins")
        inserted.set(qn("w:id"), "17")
        inserted.set(qn("w:author"), "Ada")
        run = OxmlElement("w:r")
        text = OxmlElement("w:t")
        text.text = "preserved revision"
        run.append(text)
        inserted.append(run)
        doc.paragraphs[0]._p.append(inserted)
        math = OxmlElement("m:oMath")
        math_run = OxmlElement("m:r")
        math_text = OxmlElement("m:t")
        math_text.text = "x"
        math_run.append(math_text)
        math.append(math_run)
        doc.paragraphs[0]._p.append(math)
        doc.save(path)
        before.update(_parts(path))

    monkeypatch.setattr(engine, "postprocess_main", seed_rich_content)
    script = wongo_project / "finish.py"
    script.write_text(
        "import os\nfrom docx import Document\n"
        "for path in os.environ['WONGO_POSTPROCESS_OUTPUTS'].split(os.pathsep):\n"
        "    doc = Document(path)\n"
        "    for element in doc.element.iter():\n"
        "        if element.tag.endswith('}docPr'):\n"
        "            element.set('descr', 'Verified schematic')\n"
        "    doc.save(path)\n",
        encoding="utf-8",
    )
    config = wongo_project / "_journal.yml"
    config.write_text(config.read_text(encoding="utf-8")
                      + "postprocess: [finish.py]\n", encoding="utf-8")
    result = render_project(wongo_project, "collab")
    output = result.outputs[0]
    after = _parts(output)
    assert b"Verified schematic" in after["word/document.xml"]
    assert b"preserved revision" in after["word/document.xml"]
    assert b"<m:oMath>" in after["word/document.xml"]
    assert before.keys() == after.keys()
    assert all(data == after[name] for name, data in before.items()
               if name != "word/document.xml")
    expected = etree.fromstring(before["word/document.xml"])
    for element in expected.iter():
        if element.tag.endswith("}docPr"):
            element.set("descr", "Verified schematic")
    actual = etree.fromstring(after["word/document.xml"])
    assert etree.tostring(expected, method="c14n") == etree.tostring(actual, method="c14n")
    assert output_freshness(wongo_project, "collab")["state"] == "fresh"
    entry = read_manifest(wongo_project)["targets"]["collab"]
    assert entry["outputs"][output.name] == entry["postprocess"][0]["outputs"][output.name]["after"]

    repeated = tmp_path / "repeated.docx"
    repeated.write_bytes(output.read_bytes())
    env = os.environ.copy()
    env["WONGO_POSTPROCESS_OUTPUTS"] = str(repeated)
    subprocess.run([sys.executable, str(script)], env=env, check=True)
    assert _parts(repeated) == after


def test_transform_reordering_invalidates_the_build(stub_quarto, wongo_project):
    config = wongo_project / "_journal.yml"
    original = config.read_text(encoding="utf-8")
    for name in ("first.py", "second.py"):
        (wongo_project / name).write_text("# identity transform\n", encoding="utf-8")
    config.write_text(original + "postprocess: [first.py, second.py]\n", encoding="utf-8")
    render_project(wongo_project, "collab")
    assert output_freshness(wongo_project, "collab")["state"] == "fresh"
    config.write_text(original + "postprocess: [second.py, first.py]\n", encoding="utf-8")
    assert output_freshness(wongo_project, "collab")["state"] == "stale"
